"""Spreadsheet format readers supporting modern XLSX/XLSM, legacy binary XLS, and XML Spreadsheet 2003."""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from typing import Any
from zipfile import BadZipFile

import charset_normalizer
import openpyxl
import python_calamine
import xlrd

from sarathi.sankalpa import (
    CanonicalDocument,
    ProvenanceRecord,
    TableData,
    WarningRecord,
)
from sarathi.shakti.native_extraction.readers.common import (
    CAPABILITY_ID,
    PLUGIN_ID,
    STAGE_NAME,
)


def _sheet_sort_key(fname: str) -> int:
    """Extract digits from worksheet filename for natural ordering."""
    digits = "".join(ch for ch in fname if ch.isdigit())
    return int(digits) if digits else 0


def read_xlsx(
    data: bytes,
    input_id: str,
) -> tuple[CanonicalDocument, tuple[ProvenanceRecord, ...], tuple[WarningRecord, ...]]:
    """Extract all sheets from XLSX/XLSM using python-calamine with openpyxl fallback."""
    tables: list[TableData] = []
    provenances: list[ProvenanceRecord] = []
    warnings: list[WarningRecord] = []
    reader_used = "python-calamine"

    try:
        wb = python_calamine.CalamineWorkbook.from_object(io.BytesIO(data))
        for sheet_name in wb.sheet_names:
            sheet = wb.get_sheet_by_name(sheet_name)
            raw_rows = sheet.to_python()
            if raw_rows and len(raw_rows) > 0:
                headers = tuple(str(col or "") for col in raw_rows[0])
                data_rows = tuple(tuple(cell for cell in row) for row in raw_rows[1:])
            else:
                headers = ()
                data_rows = ()

            tables.append(
                TableData(
                    name=sheet_name,
                    headers=headers,
                    rows=data_rows,
                )
            )
            provenances.append(
                ProvenanceRecord(
                    source_input_id=input_id,
                    stage=STAGE_NAME,
                    plugin_id=PLUGIN_ID,
                    capability_id=CAPABILITY_ID,
                    evidence={"reader": "python-calamine", "sheet": sheet_name, "row_count": len(raw_rows)},
                )
            )
    except (python_calamine.CalamineError, BadZipFile):
        # Fallback to openpyxl
        reader_used = "openpyxl"
        warnings.append(
            WarningRecord(
                code="CALAMINE_FALLBACK",
                message="python-calamine failed on workbook, fell back to openpyxl reader.",
                stage=STAGE_NAME,
            )
        )
        wb_px = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            for ws in wb_px.worksheets:
                rows_list = list(ws.iter_rows(values_only=True))
                if rows_list and len(rows_list) > 0:
                    headers = tuple(str(c or "") for c in rows_list[0])
                    data_rows = tuple(tuple(c for c in row) for row in rows_list[1:])
                else:
                    headers = ()
                    data_rows = ()

                tables.append(
                    TableData(
                        name=ws.title,
                        headers=headers,
                        rows=data_rows,
                    )
                )
                provenances.append(
                    ProvenanceRecord(
                        source_input_id=input_id,
                        stage=STAGE_NAME,
                        plugin_id=PLUGIN_ID,
                        capability_id=CAPABILITY_ID,
                        evidence={"reader": "openpyxl", "sheet": ws.title, "row_count": len(rows_list)},
                    )
                )
        finally:
            wb_px.close()

    # Check for active filters and hidden rows across worksheets via direct XML inspection
    filter_detected = False
    filter_details: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            sheet_titles: list[str] = []
            if "xl/workbook.xml" in zf.namelist():
                wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
                for elem in wb_root.iter():
                    if elem.tag.endswith("sheet") and "name" in elem.attrib:
                        sheet_titles.append(elem.attrib["name"])

            sheet_files = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
            sheet_files.sort(key=_sheet_sort_key)

            for idx, s_name in enumerate(sheet_files):
                sheet_bytes = zf.read(s_name)
                has_af = b"<autoFilter" in sheet_bytes
                has_hidden = b'hidden="1"' in sheet_bytes or b'hidden="true"' in sheet_bytes
                if has_af or has_hidden:
                    filter_detected = True
                    s_tree = ET.fromstring(sheet_bytes)
                    af_elem = next((e for e in s_tree.iter() if e.tag.endswith("autoFilter")), None)
                    af_ref = af_elem.attrib.get("ref", "none") if af_elem is not None else "none"
                    title = sheet_titles[idx] if idx < len(sheet_titles) else f"Sheet{idx+1}"
                    filter_details.append(f"{title}(auto_filter={af_ref}, hidden_rows={has_hidden})")
                    warnings.append(
                        WarningRecord(
                            code="EXCEL_FILTER_APPLIED",
                            message=f"Worksheet '{title}' has an AutoFilter ({af_ref}) or hidden rows. All rows are extracted.",
                            stage=STAGE_NAME,
                        )
                    )
    except Exception as exc:
        warnings.append(
            WarningRecord(
                code="EXCEL_FILTER_CHECK_SKIPPED",
                message=f"Worksheet filter inspection could not be completed: {type(exc).__name__}.",
                stage=STAGE_NAME,
            )
        )

    # Build composite text from table headers and cells for downstream processing
    composite_lines: list[str] = []
    for t in tables:
        if t.name:
            composite_lines.append(f"--- Sheet: {t.name} ---")
        if t.headers:
            composite_lines.append(" | ".join(str(h) for h in t.headers if str(h).strip()))
        for row in t.rows:
            row_str = " | ".join(str(c) for c in row if c is not None and str(c).strip())
            if row_str:
                composite_lines.append(row_str)
    full_text = "\n".join(composite_lines)

    meta: dict[str, Any] = {"reader": reader_used}
    if filter_detected:
        meta["filter_applied"] = True
        meta["filter_details"] = "; ".join(filter_details)

    canonical_doc = CanonicalDocument(
        document_id=f"doc-{input_id}",
        source_input_id=input_id,
        text=full_text,
        tables=tuple(tables),
        detected_type="xlsx",
        metadata=meta,
    )
    return canonical_doc, tuple(provenances), tuple(warnings)


def read_xls_legacy(
    data: bytes,
    input_id: str,
) -> tuple[CanonicalDocument, tuple[ProvenanceRecord, ...], tuple[WarningRecord, ...]]:
    """Extract all sheets from legacy BIFF .xls using xlrd."""
    tables: list[TableData] = []
    provenances: list[ProvenanceRecord] = []
    warnings: list[WarningRecord] = []
    composite_lines: list[str] = []

    rb = xlrd.open_workbook(file_contents=data)
    for sheet_idx in range(rb.nsheets):
        sheet = rb.sheet_by_index(sheet_idx)
        rows_list: list[tuple[Any, ...]] = [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
        if rows_list and len(rows_list) > 0:
            headers = tuple(str(c or "") for c in rows_list[0])
            data_rows = tuple(rows_list[1:])
        else:
            headers = ()
            data_rows = ()

        tables.append(
            TableData(
                name=sheet.name,
                headers=headers,
                rows=data_rows,
            )
        )
        provenances.append(
            ProvenanceRecord(
                source_input_id=input_id,
                stage=STAGE_NAME,
                plugin_id=PLUGIN_ID,
                capability_id=CAPABILITY_ID,
                evidence={"reader": "xlrd", "sheet": sheet.name, "row_count": len(rows_list)},
            )
        )
        composite_lines.append(f"--- Sheet: {sheet.name} ---")
        composite_lines.append(" | ".join(headers))
        for row in data_rows:
            composite_lines.append(" | ".join(str(c) for c in row))

    canonical_doc = CanonicalDocument(
        document_id=f"doc-{input_id}",
        source_input_id=input_id,
        text="\n".join(composite_lines),
        tables=tuple(tables),
        detected_type="xls_legacy",
    )
    return canonical_doc, tuple(provenances), tuple(warnings)


def read_spreadsheet_ml(
    data: bytes,
    input_id: str,
) -> tuple[CanonicalDocument, tuple[ProvenanceRecord, ...], tuple[WarningRecord, ...]]:
    """Extract XML spreadsheets (2003 XML format) using ElementTree."""
    tables: list[TableData] = []
    provenances: list[ProvenanceRecord] = []
    warnings: list[WarningRecord] = []

    try:
        root = ET.fromstring(data)
    except (ET.ParseError, UnicodeDecodeError):
        match = charset_normalizer.from_bytes(data).best()
        enc = match.encoding if match and match.encoding else "utf-8"
        text_content = data.decode(enc, errors="replace")
        if text_content.startswith("<?xml"):
            end_decl = text_content.find("?>")
            if end_decl != -1:
                text_content = text_content[end_decl + 2 :].lstrip()
        root = ET.fromstring(text_content.encode("utf-8"))

    # Strip namespace for robust tag matching
    ns = ""
    if root.tag.startswith("{"):
        ns = root.tag.split("}")[0] + "}"

    worksheets = root.findall(f"{ns}Worksheet")
    for ws_idx, ws in enumerate(worksheets, 1):
        ws_name = ws.attrib.get(f"{ns}Name", f"Sheet_{ws_idx}")
        table_elem = ws.find(f"{ns}Table")
        rows: list[tuple[Any, ...]] = []
        headers: tuple[str, ...] = ()

        if table_elem is not None:
            for r_idx, row_elem in enumerate(table_elem.findall(f"{ns}Row")):
                cell_vals: list[str] = []
                for cell_elem in row_elem.findall(f"{ns}Cell"):
                    idx_str = cell_elem.attrib.get(f"{ns}Index") or cell_elem.attrib.get("Index")
                    if idx_str is not None:
                        try:
                            target_col = int(idx_str) - 1
                            while len(cell_vals) < target_col:
                                cell_vals.append("")
                        except ValueError:
                            pass
                    data_elem = cell_elem.find(f"{ns}Data")
                    val = data_elem.text if data_elem is not None and data_elem.text else ""
                    cell_vals.append(val.strip())

                if r_idx == 0:
                    headers = tuple(cell_vals)
                else:
                    rows.append(tuple(cell_vals))

        tables.append(
            TableData(
                name=ws_name,
                headers=headers,
                rows=tuple(rows),
            )
        )
        provenances.append(
            ProvenanceRecord(
                source_input_id=input_id,
                stage=STAGE_NAME,
                plugin_id=PLUGIN_ID,
                capability_id=CAPABILITY_ID,
                evidence={"reader": "elementtree", "sheet": ws_name, "row_count": len(rows)},
            )
        )

    composite_lines: list[str] = []
    for t in tables:
        composite_lines.append(" ".join(t.headers))
        for row in t.rows:
            composite_lines.append(" ".join(str(cell) for cell in row if cell is not None and str(cell).strip()))
    composite_text = "\n".join(composite_lines)

    canonical_doc = CanonicalDocument(
        document_id=f"doc-{input_id}",
        source_input_id=input_id,
        text=composite_text,
        tables=tuple(tables),
        detected_type="spreadsheet_ml",
    )
    return canonical_doc, tuple(provenances), tuple(warnings)
