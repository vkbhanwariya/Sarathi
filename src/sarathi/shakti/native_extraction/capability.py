"""Shruti - Native Extraction Executable Capability."""

from __future__ import annotations

import csv
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Callable
from zipfile import BadZipFile

import openpyxl.utils.exceptions
import pymupdf
import python_calamine
import xlrd

from sarathi.dosh import DoshError, FailureCode
from sarathi.sankalpa import (
    ArtifactIntent,
    ArtifactPayload,
    CanonicalDocument,
    CapabilityDeclaration,
    ExecutionContext,
    ProvenanceRecord,
    Request,
    Result,
    WarningRecord,
)
from sarathi.shakti.docx_exporter import build_docx_payload
from sarathi.shakti.native_extraction.detector import DetectedFormat, detect_content_format
from sarathi.shakti.native_extraction.plugin import CAPABILITY_DECLARATION
from sarathi.shakti.native_extraction.readers import (
    read_csv_or_text,
    read_docx,
    read_html_table,
    read_pdf,
    read_spreadsheet_ml,
    read_xls_legacy,
    read_xlsx,
)

_PARSE_EXCEPTIONS = (
    pymupdf.FileDataError,
    pymupdf.EmptyFileError,
    openpyxl.utils.exceptions.InvalidFileException,
    BadZipFile,
    python_calamine.CalamineError,
    xlrd.biffh.XLRDError,
    ET.ParseError,
    csv.Error,
    UnicodeDecodeError,
)


def _get_reader(
    fmt: DetectedFormat,
) -> Callable[[bytes, str], tuple[CanonicalDocument, list[ProvenanceRecord], list[WarningRecord]]] | None:
    """Return the concrete reader for a detected format using modern structural pattern matching."""
    match fmt:
        case DetectedFormat.PDF:
            return read_pdf
        case DetectedFormat.DOCX:
            return read_docx
        case DetectedFormat.XLSX:
            return read_xlsx
        case DetectedFormat.XLS_LEGACY:
            return read_xls_legacy
        case DetectedFormat.HTML_TABLE:
            return read_html_table
        case DetectedFormat.SPREADSHEET_ML:
            return read_spreadsheet_ml
        case DetectedFormat.CSV_OR_TEXT:
            return read_csv_or_text
        case _:
            return None


def _has_usable_content(doc: CanonicalDocument) -> bool:
    """Check whether a CanonicalDocument contains usable text or table data."""
    has_text = bool(doc.text and doc.text.strip()) or any(bool(p.text and p.text.strip()) for p in doc.pages)
    has_tables = any(len(t.rows) > 0 or len(t.headers) > 0 for t in doc.tables) or any(
        any(len(t.rows) > 0 or len(t.headers) > 0 for t in p.tables) for p in doc.pages
    )
    return has_text or has_tables


class NativeExtractionCapability:
    """Canonical executable capability for Shruti Native Extraction."""

    def __init__(self, declaration: CapabilityDeclaration = CAPABILITY_DECLARATION) -> None:
        self.declaration: CapabilityDeclaration = declaration

    def execute(
        self,
        request: Request,
        context: ExecutionContext,
        prior_result: Result | None = None,
    ) -> Result:
        """Execute byte-first native extraction across request inputs.

        Returns canonical Document/Table data on success, escalates to OCR if native
        content is empty or unreadable, and raises DoshError(FailureCode.UNSUPPORTED)
        on unsupported binary content.
        """
        if not isinstance(request, Request):
            raise TypeError(f"request must be a Request instance, got {type(request).__name__}.")
        if not isinstance(context, ExecutionContext):
            raise TypeError(f"context must be an ExecutionContext instance, got {type(context).__name__}.")
        if prior_result is not None and not isinstance(prior_result, Result):
            raise TypeError(f"prior_result must be a Result instance or None, got {type(prior_result).__name__}.")

        extracted_docs: list[CanonicalDocument] = []
        all_provenance: list[ProvenanceRecord] = []
        all_warnings: list[WarningRecord] = []
        needs_ocr = False

        for inp in request.inputs:
            # Read input file bytes
            try:
                data = inp.source_path.read_bytes()
            except OSError as exc:
                # File I/O error or missing file
                raise DoshError(
                    code=FailureCode.EXECUTION_FAILED,
                    message="Failed to read source input file.",
                ) from exc

            # Detect format by inspecting content bytes
            fmt = detect_content_format(data, inp.source_path)

            if fmt == DetectedFormat.UNKNOWN:
                # If file is empty, escalate to OCR
                if len(data) == 0:
                    needs_ocr = True
                    all_warnings.append(
                        WarningRecord(
                            code="EMPTY_INPUT",
                            message="Input file is empty. Escalating to OCR.",
                            stage="read_native",
                        )
                    )
                    empty_doc = CanonicalDocument(
                        document_id=f"doc-{inp.input_id}",
                        source_input_id=inp.input_id,
                    )
                    extracted_docs.append(empty_doc)
                    continue

                # Genuinely unsupported binary format
                raise DoshError(
                    code=FailureCode.UNSUPPORTED,
                    message="Unsupported content format for native extraction.",
                )

            reader = _get_reader(fmt)
            if reader is None:
                raise DoshError(
                    code=FailureCode.UNSUPPORTED,
                    message="Unsupported content format for native extraction.",
                )

            # Route to concrete native readers with honest parse error handling
            try:
                doc, provs, warns = reader(data, inp.input_id)
                extracted_docs.append(doc)
                all_provenance.extend(provs)
                all_warnings.extend(warns)

                if not _has_usable_content(doc):
                    # Empty native content (e.g. scanned PDF with no text stream) -> escalate to OCR
                    needs_ocr = True
                    all_warnings.append(
                        WarningRecord(
                            code="NATIVE_EXTRACTION_EMPTY",
                            message="No usable native text or tables found in document. Escalating to OCR.",
                            stage="read_native",
                        )
                    )

            except DoshError:
                raise
            except _PARSE_EXCEPTIONS:
                # Corrupted or unparseable document -> escalate to OCR
                needs_ocr = True
                all_warnings.append(
                    WarningRecord(
                        code="NATIVE_PARSE_ERROR",
                        message="Failed to parse document content natively. Escalating to OCR.",
                        stage="read_native",
                    )
                )
                corrupt_doc = CanonicalDocument(
                    document_id=f"doc-{inp.input_id}",
                    source_input_id=inp.input_id,
                )
                extracted_docs.append(corrupt_doc)

        result_data: Any = extracted_docs[0] if len(extracted_docs) == 1 else tuple(extracted_docs)

        payloads: list[ArtifactPayload] = []
        if not needs_ocr:
            for inp, doc in zip(request.inputs, extracted_docs):
                has_content = bool(doc.text.strip()) or bool(doc.tables) or any(p.text.strip() or p.tables for p in doc.pages)
                if has_content:
                    stem = Path(inp.display_name).stem if inp.display_name else inp.input_id
                    if doc.text.strip():
                        if len(doc.pages) > 1:
                            page_sections = []
                            for p in doc.pages:
                                heading = f"--- Page {p.page_number} ---"
                                if p.text:
                                    page_sections.append(f"{heading}\n{p.text}")
                                else:
                                    page_sections.append(heading)
                            txt_content = "\n\n".join(page_sections)
                        else:
                            txt_content = doc.text
                    elif doc.tables:
                        table_lines = []
                        for t in doc.tables:
                            if t.headers:
                                table_lines.append(" | ".join(str(c) for c in t.headers))
                            for r in t.rows:
                                table_lines.append(" | ".join(str(c) for c in r))
                        txt_content = "\n".join(table_lines)
                    else:
                        txt_content = ""
                    payloads.append(
                        ArtifactPayload(
                            intent=ArtifactIntent(
                                name=f"{stem}_extracted.txt",
                                role="extracted_text",
                                media_type="text/plain",
                            ),
                            content=txt_content.encode("utf-8"),
                        )
                    )
                    payloads.append(
                        build_docx_payload(
                            doc=doc,
                            filename=f"{stem}_extracted.docx",
                            role="extracted_document",
                        )
                    )

        return Result(
            data=result_data,
            artifact_payloads=tuple(payloads),
            warnings=tuple(all_warnings),
            provenance=tuple(all_provenance),
            next_requirement="ocr" if needs_ocr else None,
        )
