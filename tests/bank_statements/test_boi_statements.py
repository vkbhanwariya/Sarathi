"""Comprehensive Unit & Integration Tests for Bank of India Statements.

Validates:
1. BOI profile identification and in-table identity extraction (ACC_NO, ACCT_NAME, ACCT_TYPE).
2. Universal rule: Check that credit and debit cannot both be blank or zero (BOTH_AMOUNTS_BLANK).
3. 26 transactions extraction with zero-normalization (0 -> None in debit/credit).
4. Financial balance reconciliation: Credits - Debits == Final Balance.
5. Option B canonical deliverables and multi-sheet List_of_Accounts.xlsx directory.
"""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from sarathi.sankalpa import CanonicalDocument, ExecutionContext, InputRef, Request, Result, TableData
from sarathi.shakti.bank_statements.capability import BankStatementCapability
from sarathi.shakti.bank_statements.detector import detect_bank_statement
from sarathi.shakti.bank_statements.mapper import load_bank_profile_yaml
from sarathi.shakti.bank_statements.models import ValidationStatus
from sarathi.shakti.native_extraction.readers.spreadsheet import read_xlsx

_BOI_FILE = Path("Input/Bank of India/663618210000712.xlsx")


@pytest.mark.skipif(not _BOI_FILE.exists(), reason="Input/Bank of India statement not present")
def test_boi_statement_detection_and_extraction() -> None:
    data = _BOI_FILE.read_bytes()
    doc, _, _ = read_xlsx(data, _BOI_FILE.name)

    det = detect_bank_statement(doc)
    assert det.is_bank_statement is True
    assert det.bank_name == "Bank of India"

    cap = BankStatementCapability()
    in_ref = InputRef(
        input_id="in-boi-01",
        source_path=_BOI_FILE,
        display_name=_BOI_FILE.name,
        size_bytes=len(data),
    )
    req = Request(request_id="req-boi-01", requirement="bank_statement", inputs=(in_ref,))
    ctx = ExecutionContext(run_id="run-boi-01", request_id="req-boi-01", trace_id="tr-1", span_id="sp-1")
    res = cap.execute(req, ctx, Result(data=doc))

    consolidation = res.data
    assert consolidation.total_transactions == 26
    assert len(consolidation.statements) == 1

    stmt = consolidation.statements[0]
    assert stmt.bank_name == "Bank of India"
    assert stmt.account_identity is not None
    assert stmt.account_identity.account_holder == "MOHAMMAD SARIF"
    assert stmt.account_identity.masked_account_number.endswith("0712")
    assert stmt.account_identity.account_type == "SAVING"
    assert len(stmt.transactions) == 26
    assert len(stmt.issues) == 0

    # Financial balance check: 46750 (credit) - 40000 (debit) = 6750
    assert consolidation.total_credit == Decimal("46750")
    assert consolidation.total_debit == Decimal("40000")
    net_diff = consolidation.total_credit - consolidation.total_debit
    assert net_diff == stmt.transactions[-1].running_balance


@pytest.mark.skipif(not _BOI_FILE.exists(), reason="Input/Bank of India statement not present")
def test_boi_deliverables_and_formatting() -> None:
    data = _BOI_FILE.read_bytes()
    doc, _, _ = read_xlsx(data, _BOI_FILE.name)

    cap = BankStatementCapability()
    in_ref = InputRef(
        input_id="in-boi-02",
        source_path=_BOI_FILE,
        display_name=_BOI_FILE.name,
        size_bytes=len(data),
    )
    req = Request(request_id="req-boi-02", requirement="bank_statement", inputs=(in_ref,))
    ctx = ExecutionContext(run_id="run-boi-02", request_id="req-boi-02", trace_id="tr-2", span_id="sp-2")
    res = cap.execute(req, ctx, Result(data=doc))
    assert len(res.artifact_payloads) == 3

    # 1. List_of_Accounts.xlsx
    acc_payload = next(a for a in res.artifact_payloads if a.intent.name == "List_of_Accounts.xlsx")
    wb_acc = openpyxl.load_workbook(io.BytesIO(acc_payload.content))
    assert wb_acc.sheetnames == ["Accounts", "Processing_Summary"]

    ws_acc = wb_acc["Accounts"]
    acc_rows = list(ws_acc.iter_rows(values_only=True))
    assert acc_rows[1][1] == "MOHAMMAD SARIF"
    assert "0712" in acc_rows[1][2]
    assert acc_rows[1][3] == "Bank of India"

    ws_sum = wb_acc["Processing_Summary"]
    sum_rows = list(ws_sum.iter_rows(values_only=True))
    assert sum_rows[1][1] == _BOI_FILE.stem
    assert sum_rows[1][6] == 26  # Total rows scanned
    assert sum_rows[1][7] == 26  # Successful transactions
    assert sum_rows[1][8] == 0   # Duplicates
    assert sum_rows[1][9] == 0   # Failed/skipped rows
    assert sum_rows[1][10] == "VALID"
    assert sum_rows[1][12] == "ALL_CLEAR"

    # 2. Consolidated_Transactions.xlsx
    txn_payload = next(a for a in res.artifact_payloads if a.intent.name == "Consolidated_Transactions.xlsx")
    wb_txn = openpyxl.load_workbook(io.BytesIO(txn_payload.content))
    ws_txn = wb_txn.active
    # First row is credit 200, debit must be None (not 0.00)
    assert ws_txn.cell(row=2, column=7).value is None
    assert ws_txn.cell(row=2, column=8).value == 200
    # Fifth row is debit 5000, credit must be None (not 0.00)
    assert ws_txn.cell(row=6, column=7).value == 5000
    assert ws_txn.cell(row=6, column=8).value is None

    total_row = len(res.data.transactions) + 2
    assert ws_txn.cell(row=total_row, column=7).value == f"=SUBTOTAL(9, G2:G{total_row-1})"
    assert ws_txn.cell(row=total_row, column=8).value == f"=SUBTOTAL(9, H2:H{total_row-1})"

    # 3. Consolidated_Bank_Statement.parquet
    pq_payload = next(a for a in res.artifact_payloads if a.intent.name == "Consolidated_Bank_Statement.parquet")
    df = pl.read_parquet(io.BytesIO(pq_payload.content))
    assert len(df) == 26
    assert "transaction_hash" in df.columns
    assert "transaction_mode" in df.columns


def test_both_debit_and_credit_blank_check() -> None:
    """Universal rule: If both debit and credit are blank/empty, row is marked INVALID with BOTH_AMOUNTS_BLANK."""
    cap = BankStatementCapability()
    table = TableData(
        name="test_table",
        headers=("Date", "Description", "Debit", "Credit", "Balance"),
        rows=(
            ("01-01-2026", "Valid Credit", "", "500", "1500"),
            ("02-01-2026", "Blank Row", "", "", "1500"),
            ("03-01-2026", "Valid Debit", "200", "", "1300"),
        ),
    )
    doc = CanonicalDocument(
        document_id="doc-blank-test",
        source_input_id="inp-blank-test",
        text="Bank Statement Bank of India",
        tables=(table,),
    )
    ctx = ExecutionContext(run_id="r-chk-1", request_id="req-chk-1", trace_id="t1", span_id="s1")
    req = Request(
        request_id="req-chk-1",
        requirement="bank_statements",
        inputs=(InputRef("inp-blank-test", Path("test.xlsx"), "test.xlsx", 100),),
    )
    res = cap.execute(req, ctx, prior_result=Result(data=doc, provenance=()))

    stmt = res.data.statements[0]
    # Blank row has BOTH_AMOUNTS_BLANK issue
    blank_tx = stmt.transactions[1]
    assert blank_tx.status == ValidationStatus.INVALID
    assert any(i.code == "BOTH_AMOUNTS_BLANK" for i in blank_tx.issues)

    # Valid consolidated transactions must only contain the 2 valid transactions
    assert len(res.data.transactions) == 2
    assert res.data.total_credit == Decimal("500")
    assert res.data.total_debit == Decimal("200")


def test_both_debit_and_credit_zero_check() -> None:
    """Universal rule: If both debit and credit are explicitly 0, row is marked INVALID with BOTH_AMOUNTS_BLANK."""
    cap = BankStatementCapability()
    table = TableData(
        name="test_table_zero",
        headers=("Date", "Description", "Debit", "Credit", "Balance"),
        rows=(
            ("01-01-2026", "Both Zero Row", "0", "0", "1000"),
            ("02-01-2026", "Valid Debit", "100", "0", "900"),
        ),
    )
    doc = CanonicalDocument(
        document_id="doc-zero-test",
        source_input_id="inp-zero-test",
        text="Bank Statement Bank of India",
        tables=(table,),
    )
    ctx = ExecutionContext(run_id="r-chk-2", request_id="req-chk-2", trace_id="t2", span_id="s2")
    req = Request(
        request_id="req-chk-2",
        requirement="bank_statements",
        inputs=(InputRef("inp-zero-test", Path("test.xlsx"), "test.xlsx", 100),),
    )
    res = cap.execute(req, ctx, prior_result=Result(data=doc, provenance=()))

    stmt = res.data.statements[0]
    zero_tx = stmt.transactions[0]
    assert zero_tx.status == ValidationStatus.INVALID
    assert any(i.code == "BOTH_AMOUNTS_BLANK" for i in zero_tx.issues)

    # Valid consolidated transactions must only contain the 1 valid transaction
    assert len(res.data.transactions) == 1
    assert res.data.total_debit == Decimal("100")


def test_boi_profile_file_format_metadata() -> None:
    prof = load_bank_profile_yaml(Path("src/sarathi/data/banks/boi_excel_fmt1.yaml"))
    assert prof["profile_id"] == "boi_excel_fmt1"
    assert prof["bank_name"] == "Bank of India"
    assert "file_format" in prof
    ff = prof["file_format"]
    assert ff["declared_extension"] == ".xlsx"
    assert "openxml" in ff["actual_hidden_format"].lower()
    assert ff["mime_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
