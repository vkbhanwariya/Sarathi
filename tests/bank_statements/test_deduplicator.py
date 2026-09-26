"""Consolidated unit and invariant tests for bank statement deduplication and ordering."""

from __future__ import annotations

import io
import time as time_mod
from datetime import date, time
from decimal import Decimal
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from sarathi.sankalpa import (
    CanonicalDocument,
    InputRef,
    PageData,
    ProvenanceRecord,
    Request,
    TableData,
)
from sarathi.shakti.bank_statements.capability import BankStatementCapability
from sarathi.shakti.bank_statements.consolidator import (
    build_accounts_xlsx_artifact,
    build_parquet_artifact,
    consolidate_statements,
)
from sarathi.shakti.bank_statements.deduplicator import deduplicate_transactions
from sarathi.shakti.bank_statements.models import (
    BankStatement,
    DuplicateDecision,
    Transaction,
    ValidationStatus,
    create_account_identity,
)
from sarathi.shakti.bank_statements.validator import validate_statement_balances


def test_deduplicate_exact_transactions() -> None:
    """Exact match on date, amount, description, ref, and balance drops the duplicate."""
    ident = create_account_identity("State Bank of India", "30123456789")
    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="UPI Payment",
        bank_name="State Bank of India",
        account_identity=ident,
        debit=Decimal("100.00"),
        running_balance=Decimal("5000.00"),
        reference_number="REF123",
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="UPI Payment",
        bank_name="State Bank of India",
        account_identity=ident,
        debit=Decimal("100.00"),
        running_balance=Decimal("5000.00"),
        reference_number="REF123",
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 1
    assert len(res.duplicates) == 1
    assert res.duplicates[0][2] == DuplicateDecision.PROVEN_DUPLICATE


def test_deduplicate_distinct_documents_without_account_identity() -> None:
    """Transactions with identical fields from distinct unlabelled documents must NOT falsely collide."""
    p1 = ProvenanceRecord(source_input_id="doc-A", capability_id="bank_statements", stage="extraction")
    p2 = ProvenanceRecord(source_input_id="doc-B", capability_id="bank_statements", stage="extraction")

    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Cash Withdrawal",
        bank_name="Unknown Bank",
        account_identity=None,
        debit=Decimal("500.00"),
        provenance=(p1,),
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Cash Withdrawal",
        bank_name="Unknown Bank",
        account_identity=None,
        debit=Decimal("500.00"),
        provenance=(p2,),
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 2
    assert len(res.duplicates) == 0


def test_deduplicate_merges_provenance_on_drop() -> None:
    """When a duplicate is dropped, its source provenance is merged into the surviving transaction."""
    ident = create_account_identity("HDFC Bank", "5010099999")
    p1 = ProvenanceRecord(source_input_id="doc-1", capability_id="bank_statements", stage="extraction", page_number=1)
    p2 = ProvenanceRecord(source_input_id="doc-1", capability_id="bank_statements", stage="extraction", page_number=2)

    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Salary Credit",
        bank_name="HDFC Bank",
        account_identity=ident,
        reference_number="SAL20260101",
        credit=Decimal("75000.00"),
        provenance=(p1,),
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Salary Credit",
        bank_name="HDFC Bank",
        account_identity=ident,
        reference_number="SAL20260101",
        credit=Decimal("75000.00"),
        provenance=(p2,),
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 1
    assert len(res.duplicates) == 1
    assert res.duplicates[0][2] == DuplicateDecision.PROVEN_DUPLICATE
    surviving = res.unique_transactions[0]
    assert len(surviving.provenance) == 2
    assert p1 in surviving.provenance
    assert p2 in surviving.provenance


def test_deduplicate_probable_retains_both_transactions() -> None:
    """Without reference or running balance, matching date/amount/narration is PROBABLE and retains both."""
    ident = create_account_identity("HDFC Bank", "5010099999")
    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="ATM Cash Withdrawal",
        bank_name="HDFC Bank",
        account_identity=ident,
        debit=Decimal("5000.00"),
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="ATM Cash Withdrawal",
        bank_name="HDFC Bank",
        account_identity=ident,
        debit=Decimal("5000.00"),
    )

    res = deduplicate_transactions([tx1, tx2])
    # Both transactions are preserved per Bank Veda
    assert len(res.unique_transactions) == 2
    assert len(res.duplicates) == 1
    assert res.duplicates[0][2] == DuplicateDecision.PROBABLE_DUPLICATE
    # Second transaction receives warning issue
    assert any(iss.code == "PROBABLE_DUPLICATE_TRANSACTION" for iss in res.unique_transactions[1].issues)


def test_deduplicate_contradictory_time_retains_both_transactions() -> None:
    """Transactions with matching date/account/amount/ref but contradictory explicit times must NOT be merged."""
    ident = create_account_identity("HDFC Bank", "5010099999")
    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        transaction_time=time(9, 0, 0),
        description="Payment",
        bank_name="HDFC Bank",
        account_identity=ident,
        reference_number="REF123",
        debit=Decimal("100.00"),
        running_balance=Decimal("900.00"),
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 1),
        transaction_time=time(11, 0, 0),
        description="Payment",
        bank_name="HDFC Bank",
        account_identity=ident,
        reference_number="REF123",
        debit=Decimal("100.00"),
        running_balance=Decimal("900.00"),
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 2
    assert len(res.duplicates) == 0


def test_continuous_monotonic_sequence_id_across_tables() -> None:
    """Verify sequence_id increments monotonically across multiple tables/pages instead of resetting."""
    cap = BankStatementCapability()

    table1 = TableData(
        headers=("Date", "Description", "Debit", "Credit", "Balance"),
        rows=(
            ("01/01/2026", "Tx 1 Page 1", "100.00", "", "1000.00"),
            ("02/01/2026", "Tx 2 Page 1", "200.00", "", "800.00"),
        ),
    )
    table2 = TableData(
        headers=("Date", "Description", "Debit", "Credit", "Balance"),
        rows=(
            ("03/01/2026", "Tx 3 Page 2", "300.00", "", "500.00"),
            ("04/01/2026", "Tx 4 Page 2", "50.00", "", "450.00"),
        ),
    )

    page1 = PageData(page_number=1, tables=(table1,))
    page2 = PageData(page_number=2, tables=(table2,))

    doc = CanonicalDocument(
        document_id="doc_multi_page",
        source_input_id="bank_doc_multi_page",
        pages=(page1, page2),
    )

    req = Request(
        request_id="req-1",
        requirement="bank_statements",
        inputs=(InputRef(input_id="in-1", source_path=Path("test.pdf"), display_name="test.pdf", size_bytes=100),),
    )
    raw_txns, _, _, _ = cap._extract_table_data(doc, req, "sbi", "State Bank of India", None)
    assert len(raw_txns) == 4
    # sequence_ids must be 1, 2, 3, 4 strictly monotonic
    seq_ids = [tx.sequence_id for tx in raw_txns]
    assert seq_ids == [1, 2, 3, 4]


def test_dedup_strong_signal_running_balance_without_reference() -> None:
    """Transactions with matching running balance and no contradictions are PROVEN_DUPLICATE."""
    ident = create_account_identity("SBI", "12345678901")
    tx1 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="Electricity Bill",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("1500.00"),
        running_balance=Decimal("25000.00"),
        sequence_id=10,
    )
    tx2 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="Electricity Bill",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("1500.00"),
        running_balance=Decimal("25000.00"),
        reference_number="REF-ELEC-99",  # tx2 has reference, tx1 lacks it
        sequence_id=25,
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 1
    assert len(res.duplicates) == 1
    assert res.duplicates[0][2] == DuplicateDecision.PROVEN_DUPLICATE
    surviving = res.unique_transactions[0]
    # Sequence id of original surviving transaction preserved
    assert surviving.sequence_id == 10
    # Reference number enriched from tx2
    assert surviving.reference_number == "REF-ELEC-99"


def test_dedup_contradictory_running_balance_rejected() -> None:
    """Transactions with identical date, amount, narration but conflicting running balance are DISTINCT."""
    ident = create_account_identity("SBI", "12345678901")
    tx1 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="ATM Cash Withdrawal",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("2000.00"),
        running_balance=Decimal("18000.00"),
        sequence_id=1,
    )
    tx2 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="ATM Cash Withdrawal",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("2000.00"),
        running_balance=Decimal("16000.00"),
        sequence_id=2,
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 2
    assert len(res.duplicates) == 0


def test_dedup_contradictory_reference_number_rejected() -> None:
    """Transactions with identical date, amount, narration but conflicting references are DISTINCT."""
    ident = create_account_identity("SBI", "12345678901")
    tx1 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="UPI Payment",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("500.00"),
        reference_number="UPI-0001",
        sequence_id=1,
    )
    tx2 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="UPI Payment",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("500.00"),
        reference_number="UPI-0002",
        sequence_id=2,
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 2
    assert len(res.duplicates) == 0


def test_overlapping_statement_period_deduplication() -> None:
    """Verify transactions present in overlapping statement ranges are deduplicated cleanly."""
    ident = create_account_identity("State Bank of India", "30123456789")

    # Statement 1 transactions (Jan 1 to Jan 20)
    tx1 = Transaction(
        transaction_date=date(2026, 1, 5),
        description="Electricity Bill",
        bank_name="State Bank of India",
        debit=Decimal("1200.00"),
        running_balance=Decimal("8800.00"),
        reference_number="TXN1001",
        account_identity=ident,
    )
    tx2 = Transaction(
        transaction_date=date(2026, 1, 15),
        description="Salary Credit",
        bank_name="State Bank of India",
        credit=Decimal("50000.00"),
        running_balance=Decimal("58800.00"),
        reference_number="SAL2026",
        account_identity=ident,
    )

    # Statement 2 transactions (Jan 10 to Jan 31) overlapping tx2 and adding tx3
    tx2_overlap = Transaction(
        transaction_date=date(2026, 1, 15),
        description="Salary Credit",
        bank_name="State Bank of India",
        credit=Decimal("50000.00"),
        running_balance=Decimal("58800.00"),
        reference_number="SAL2026",
        account_identity=ident,
    )
    tx3 = Transaction(
        transaction_date=date(2026, 1, 25),
        description="Grocery Store",
        bank_name="State Bank of India",
        debit=Decimal("2500.00"),
        running_balance=Decimal("56300.00"),
        reference_number="POS9911",
        account_identity=ident,
    )

    all_txns = [tx1, tx2, tx2_overlap, tx3]
    res = deduplicate_transactions(all_txns)

    assert len(res.unique_transactions) == 3
    assert len(res.duplicates) == 1
    dup = res.duplicates[0]
    assert dup[2] == DuplicateDecision.PROVEN_DUPLICATE
    assert dup[1].reference_number == "SAL2026"


def test_sparse_balance_continuity_derivation() -> None:
    """Verify sparse running balance rows (balance is None) continuity."""
    ident = create_account_identity("HDFC Bank", "50100234567890")

    tx1 = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Opening",
        bank_name="HDFC Bank",
        credit=Decimal("1000.00"),
        running_balance=Decimal("11000.00"),
        account_identity=ident,
    )
    # tx2 has no explicit balance in raw statement
    tx2 = Transaction(
        transaction_date=date(2026, 1, 2),
        description="Withdrawal without explicit balance",
        bank_name="HDFC Bank",
        debit=Decimal("2000.00"),
        running_balance=None,
        account_identity=ident,
    )
    # tx3 has explicit balance which matches derived 11000 - 2000 + 500 = 9500
    tx3 = Transaction(
        transaction_date=date(2026, 1, 3),
        description="Deposit",
        bank_name="HDFC Bank",
        credit=Decimal("500.00"),
        running_balance=Decimal("9500.00"),
        account_identity=ident,
    )

    stmt = BankStatement(
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident,
        opening_balance=Decimal("10000.00"),
        closing_balance=Decimal("9500.00"),
        transactions=(tx1, tx2, tx3),
    )

    validated = validate_statement_balances(stmt)
    assert len(validated.issues) == 0
    assert validated.status.value == "valid"


def test_consolidate_statements_flattens_and_sorts_chronologically() -> None:
    """consolidate_statements flattens all valid transactions and sorts chronologically."""
    ident = create_account_identity("SBI", "111122223333")

    tx_invalid = Transaction(
        transaction_date=date(2026, 1, 1),
        description="Invalid Missing Amounts",
        bank_name="SBI",
        account_identity=ident,
        status=ValidationStatus.INVALID,
    )
    tx_jan2 = Transaction(
        transaction_date=date(2026, 1, 2),
        description="Jan 2 Tx",
        bank_name="SBI",
        debit=Decimal("100.00"),
        sequence_id=1,
        account_identity=ident,
    )
    tx_jan5 = Transaction(
        transaction_date=date(2026, 1, 5),
        description="Jan 5 Tx",
        bank_name="SBI",
        credit=Decimal("500.00"),
        sequence_id=2,
        account_identity=ident,
    )
    stmt1 = BankStatement(
        bank_name="SBI",
        bank_profile="sbi",
        account_identity=ident,
        transactions=(tx_jan5, tx_jan2, tx_invalid),
    )

    tx_jan3_early = Transaction(
        transaction_date=date(2026, 1, 3),
        transaction_time=time(10, 30),
        description="Jan 3 Early Tx",
        bank_name="SBI",
        debit=Decimal("50.00"),
        sequence_id=1,
        account_identity=ident,
    )
    tx_jan3_late = Transaction(
        transaction_date=date(2026, 1, 3),
        transaction_time=time(15, 45),
        description="Jan 3 Late Tx",
        bank_name="SBI",
        credit=Decimal("300.00"),
        sequence_id=2,
        account_identity=ident,
    )
    stmt2 = BankStatement(
        bank_name="SBI",
        bank_profile="sbi",
        account_identity=ident,
        transactions=(tx_jan3_late, tx_jan3_early),
    )

    res = consolidate_statements([stmt1, stmt2])
    # tx_invalid excluded from valid transactions
    assert len(res.transactions) == 4
    assert res.transactions[0].description == "Jan 2 Tx"
    assert res.transactions[1].description == "Jan 3 Early Tx"
    assert res.transactions[2].description == "Jan 3 Late Tx"
    assert res.transactions[3].description == "Jan 5 Tx"


def test_cross_statement_deduplication_and_exports() -> None:
    """Test that transactions appearing in two overlapping statement files are deduplicated in consolidation and exports."""
    ident = create_account_identity("HDFC Bank", "5010022334455")

    # Common transaction present in both January and February statements
    overlap_tx = Transaction(
        transaction_date=date(2026, 1, 31),
        description="Salary Credit Corp",
        bank_name="HDFC Bank",
        credit=Decimal("50000.00"),
        running_balance=Decimal("65000.00"),
        account_identity=ident,
        sequence_id=10,
    )

    stmt1_tx = Transaction(
        transaction_date=date(2026, 1, 15),
        description="Groceries Store",
        bank_name="HDFC Bank",
        debit=Decimal("2500.00"),
        running_balance=Decimal("15000.00"),
        account_identity=ident,
        sequence_id=5,
    )

    stmt2_tx = Transaction(
        transaction_date=date(2026, 2, 5),
        description="Internet Bill",
        bank_name="HDFC Bank",
        debit=Decimal("1000.00"),
        running_balance=Decimal("64000.00"),
        account_identity=ident,
        sequence_id=1,
    )

    stmt1 = BankStatement(
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident,
        transactions=(stmt1_tx, overlap_tx),
        statement_id="stmt_jan",
    )

    stmt2 = BankStatement(
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident,
        transactions=(overlap_tx, stmt2_tx),
        statement_id="stmt_feb",
    )

    consolidation = consolidate_statements([stmt1, stmt2])

    # 1. Deduplication across statements: 3 unique transactions instead of 4
    assert len(consolidation.transactions) == 3
    assert consolidation.total_transactions == 3
    assert consolidation.total_credit == Decimal("50000.00")
    assert consolidation.total_debit == Decimal("3500.00")

    # Issue recorded for cross-statement duplicate
    dup_issues = [i for i in consolidation.issues if i.code == "CROSS_STATEMENT_DUPLICATE"]
    assert len(dup_issues) == 1

    # 2. Chronological ordering
    assert consolidation.transactions[0].description == "Groceries Store"
    assert consolidation.transactions[1].description == "Salary Credit Corp"
    assert consolidation.transactions[2].description == "Internet Bill"

    # 3. Parquet export strictly reflects canonical sequence
    parquet_art = build_parquet_artifact(consolidation)
    df = pl.read_parquet(io.BytesIO(parquet_art.content))
    assert len(df) == 3
    assert df["description"].to_list() == ["Groceries Store", "Salary Credit Corp", "Internet Bill"]
    assert df["credit"].to_list()[1] == Decimal("50000.00")


def test_proven_duplicate_scanned_before_probable_duplicate() -> None:
    """Item 14: Proven duplicate is matched and collapsed even when a probable duplicate was registered first."""
    import datetime

    from sarathi.shakti.bank_statements.deduplicator import deduplicate_transactions
    from sarathi.shakti.bank_statements.models import AccountIdentity, DuplicateDecision, Transaction

    ident = AccountIdentity(bank_name="Test Bank", account_fingerprint="acc-123")
    dt = datetime.date(2026, 1, 1)

    # Tx1: has no reference
    tx1 = Transaction(
        transaction_date=dt,
        description="ATM Withdrawal",
        bank_name="Test Bank",
        debit=Decimal("1000.00"),
        credit=None,
        account_identity=ident,
    )
    # Tx2: has reference 'REF100'
    tx2 = Transaction(
        transaction_date=dt,
        description="ATM Withdrawal",
        bank_name="Test Bank",
        debit=Decimal("1000.00"),
        credit=None,
        reference_number="REF100",
        account_identity=ident,
    )
    # Tx3: also has reference 'REF100' (proven duplicate of Tx2)
    tx3 = Transaction(
        transaction_date=dt,
        description="ATM Withdrawal",
        bank_name="Test Bank",
        debit=Decimal("1000.00"),
        credit=None,
        reference_number="REF100",
        account_identity=ident,
    )

    res = deduplicate_transactions([tx1, tx2, tx3])
    # Tx1 (unique), Tx2 (unique + merged with Tx3) -> exactly 2 unique transactions
    assert len(res.unique_transactions) == 2
    # The duplicate of Tx3 must be proven duplicate with Tx2, not probable duplicate with Tx1
    proven_dups = [d for d in res.duplicates if d[2] == DuplicateDecision.PROVEN_DUPLICATE]
    assert len(proven_dups) == 1
    assert proven_dups[0][0].reference_number == "REF100"
    assert proven_dups[0][1].reference_number == "REF100"


def test_duplicate_decision_no_stale_aliases() -> None:
    """Stale PROVEN and PROBABLE aliases removed from DuplicateDecision enum."""
    assert hasattr(DuplicateDecision, "PROVEN_DUPLICATE")
    assert hasattr(DuplicateDecision, "PROBABLE_DUPLICATE")
    assert hasattr(DuplicateDecision, "DISTINCT")
    assert not hasattr(DuplicateDecision, "PROVEN")
    assert not hasattr(DuplicateDecision, "PROBABLE")


def test_deduplicate_preserves_legitimate_repeated_statement_rows() -> None:
    """Verify sequence debit 100 -> credit 100 -> debit 100 with matching running balances is preserved."""
    ident = create_account_identity("HDFC Bank", "5010099999")
    prov = (ProvenanceRecord(source_input_id="doc-stmt-1", capability_id="bank_statements", stage="extraction"),)
    dt = date(2026, 1, 15)

    tx1 = Transaction(
        transaction_date=dt,
        description="ATM Withdrawal",
        bank_name="HDFC Bank",
        debit=Decimal("100.00"),
        running_balance=Decimal("900.00"),
        account_identity=ident,
        provenance=prov,
        sequence_id=1,
    )
    tx2 = Transaction(
        transaction_date=dt,
        description="UPI Deposit",
        bank_name="HDFC Bank",
        credit=Decimal("100.00"),
        running_balance=Decimal("1000.00"),
        account_identity=ident,
        provenance=prov,
        sequence_id=2,
    )
    tx3 = Transaction(
        transaction_date=dt,
        description="ATM Withdrawal",
        bank_name="HDFC Bank",
        debit=Decimal("100.00"),
        running_balance=Decimal("900.00"),
        account_identity=ident,
        provenance=prov,
        sequence_id=3,
    )

    res = deduplicate_transactions([tx1, tx2, tx3])
    # All 3 distinct rows must be preserved; total debits must remain 200
    assert len(res.unique_transactions) == 3
    total_debits = sum(t.debit for t in res.unique_transactions if t.debit is not None)
    assert total_debits == Decimal("200.00")


def test_masked_account_numbers_distinct_holders_never_deduplicate() -> None:
    """Verify that two statements displaying the same masked account number (XXXXXX1234)

    belonging to different people (Alice vs Bob) never get merged or have their transactions deleted.
    """
    ident_alice = create_account_identity("State Bank of India", "XXXXXX1234", account_holder="Alice")
    ident_bob = create_account_identity("State Bank of India", "XXXXXX1234", account_holder="Bob")

    # Both identities must have distinct fingerprints
    assert ident_alice.account_fingerprint != ident_bob.account_fingerprint

    tx_alice = Transaction(
        transaction_date=date(2026, 1, 15),
        description="Electricity Bill",
        bank_name="State Bank of India",
        debit=Decimal("500.00"),
        running_balance=Decimal("10000.00"),
        account_identity=ident_alice,
        sequence_id=1,
    )
    tx_bob = Transaction(
        transaction_date=date(2026, 1, 15),
        description="Electricity Bill",
        bank_name="State Bank of India",
        debit=Decimal("500.00"),
        running_balance=Decimal("10000.00"),
        account_identity=ident_bob,
        sequence_id=1,
    )

    stmt_alice = BankStatement(
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident_alice,
        transactions=(tx_alice,),
        statement_id="stmt_alice",
    )
    stmt_bob = BankStatement(
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident_bob,
        transactions=(tx_bob,),
        statement_id="stmt_bob",
    )

    res = consolidate_statements([stmt_alice, stmt_bob])
    # Both transactions must be preserved; neither must be deleted as a duplicate
    assert len(res.transactions) == 2
    assert res.total_debit == Decimal("1000.00")


def test_deduplication_preserves_ids_and_provenance_fields() -> None:
    """Proven duplicate merging must preserve statement_id, transaction_id, raw values, and source metadata."""
    ident = create_account_identity("State Bank of India", "30123456789")
    tx1 = Transaction(
        statement_id="stmt_sbi_month1",
        transaction_id="tx_stmt_sbi_month1_00001",
        sequence_id=1,
        transaction_date=date(2026, 1, 15),
        description="Electricity Bill Payment",
        raw_description="Electricity Bill Payment Line 1",
        raw_reference="UPI/123456/orig",
        reference_number="UPI123456orig",
        bank_name="State Bank of India",
        account_identity=ident,
        debit=Decimal("1500.00"),
        running_balance=Decimal("25000.00"),
        source_input_id="inp_month1",
        page_number=2,
        row_index=5,
    )
    tx2 = Transaction(
        statement_id="stmt_sbi_month2",
        transaction_id="tx_stmt_sbi_month2_00001",
        sequence_id=1,
        transaction_date=date(2026, 1, 15),
        description="Electricity Bill Payment",
        raw_description="Electricity Bill Payment Line 1",
        raw_reference="UPI/123456/orig",
        reference_number="UPI123456orig",
        bank_name="State Bank of India",
        account_identity=ident,
        debit=Decimal("1500.00"),
        running_balance=Decimal("25000.00"),
        source_input_id="inp_month2",
        page_number=1,
        row_index=3,
    )

    res = deduplicate_transactions([tx1, tx2])
    assert len(res.unique_transactions) == 1
    assert len(res.duplicates) == 1
    surviving = res.unique_transactions[0]
    assert surviving.statement_id == "stmt_sbi_month1"
    assert surviving.transaction_id == "tx_stmt_sbi_month1_00001"
    assert surviving.raw_description == "Electricity Bill Payment Line 1"
    assert surviving.raw_reference == "UPI/123456/orig"
    assert surviving.source_input_id == "inp_month1"
    assert surviving.page_number == 2
    assert surviving.row_index == 5


def test_consolidate_multi_account_unique_continuous_txn_ids_and_sheet_order() -> None:
    """Consolidation preserves file and sheet hierarchy (old to new) with unique continuous TXN-XXXX IDs."""
    ident_bob = create_account_identity("Bank of Baroda", "10260100019671")
    ident_au = create_account_identity("AU Small Finance Bank", "2401258671603953")

    # File 1 (BOB) - Sheet 1: 2 transactions (Jan 5 then Jan 2 in raw order)
    # File 1 (BOB) - Sheet 2: 1 transaction (Jan 1)
    tx_bob_s1_later = Transaction(
        transaction_date=date(2026, 1, 5),
        transaction_time=time(14, 0),
        description="BOB S1 Jan 5",
        bank_name="Bank of Baroda",
        account_identity=ident_bob,
        debit=Decimal("100.00"),
        source_input_id="inp_bob",
        input_location="10260100019671_S1_R05",
        page_number=1,
        row_index=5,
    )
    tx_bob_s1_earlier = Transaction(
        transaction_date=date(2026, 1, 2),
        transaction_time=time(10, 0),
        description="BOB S1 Jan 2",
        bank_name="Bank of Baroda",
        account_identity=ident_bob,
        debit=Decimal("200.00"),
        source_input_id="inp_bob",
        input_location="10260100019671_S1_R02",
        page_number=1,
        row_index=2,
    )
    tx_bob_s2 = Transaction(
        transaction_date=date(2026, 1, 1),
        transaction_time=time(9, 0),
        description="BOB S2 Jan 1",
        bank_name="Bank of Baroda",
        account_identity=ident_bob,
        credit=Decimal("500.00"),
        source_input_id="inp_bob",
        input_location="10260100019671_S2_R02",
        page_number=2,
        row_index=2,
    )
    stmt_bob = BankStatement(
        bank_name="Bank of Baroda",
        bank_profile="bob_excel_fmt1",
        account_identity=ident_bob,
        statement_id="stmt_bob",
        transactions=(tx_bob_s1_later, tx_bob_s1_earlier, tx_bob_s2),
    )

    # File 2 (AU) - Sheet 1: 2 transactions on same day with distinct times (raw order newest to oldest)
    tx_au_late = Transaction(
        transaction_date=date(2026, 2, 10),
        transaction_time=time(16, 30),
        description="AU S1 Feb 10 16:30",
        bank_name="AU Small Finance Bank",
        account_identity=ident_au,
        debit=Decimal("300.00"),
        source_input_id="inp_au",
        input_location="Statements_S1_R04",
        page_number=1,
        row_index=4,
    )
    tx_au_early = Transaction(
        transaction_date=date(2026, 2, 10),
        transaction_time=time(11, 15),
        description="AU S1 Feb 10 11:15",
        bank_name="AU Small Finance Bank",
        account_identity=ident_au,
        credit=Decimal("1000.00"),
        source_input_id="inp_au",
        input_location="Statements_S1_R05",
        page_number=1,
        row_index=5,
    )
    stmt_au = BankStatement(
        bank_name="AU Small Finance Bank",
        bank_profile="au_small_finance",
        account_identity=ident_au,
        statement_id="stmt_au",
        transactions=(tx_au_late, tx_au_early),
    )

    res = consolidate_statements([stmt_bob, stmt_au])
    assert len(res.transactions) == 5

    # 1. Continuous, unique transaction IDs
    ids = [t.transaction_id for t in res.transactions]
    assert ids == ["TXN-0001", "TXN-0002", "TXN-0003", "TXN-0004", "TXN-0005"]
    assert len(set(ids)) == 5

    # 2. Strict hierarchy:
    # First file Sheet 1 (old to new: Jan 2 -> Jan 5)
    assert res.transactions[0].description == "BOB S1 Jan 2"
    assert res.transactions[1].description == "BOB S1 Jan 5"
    # First file Sheet 2
    assert res.transactions[2].description == "BOB S2 Jan 1"
    # Second file Sheet 1 (old to new by date & time: 11:15 -> 16:30)
    assert res.transactions[3].description == "AU S1 Feb 10 11:15"
    assert res.transactions[4].description == "AU S1 Feb 10 16:30"


def test_list_of_accounts_single_row_per_account_across_multiple_statements() -> None:
    """3 monthly statements for the same account must produce exactly 1 row in the Accounts directory."""
    import io

    import openpyxl

    from sarathi.shakti.bank_statements.consolidator import (
        build_accounts_xlsx_artifact,
        consolidate_statements,
    )
    from sarathi.shakti.bank_statements.models import (
        BankStatement,
        Transaction,
        create_account_identity,
    )

    ident = create_account_identity("State Bank of India", "30123456789", account_holder="Mr. Rahul Sharma")

    tx_jan = Transaction(
        transaction_date=date(2026, 1, 10),
        description="Jan Txn",
        bank_name="State Bank of India",
        account_identity=ident,
        credit=Decimal("1000.00"),
        source_input_id="inp_jan",
        input_location="Jan_S1_R02",
        statement_id="stmt_jan",
    )
    stmt_jan = BankStatement(
        statement_id="stmt_jan",
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident,
        account_holder="Mr. Rahul Sharma",
        transactions=(tx_jan,),
        metadata={"source_file_stem": "jan_statement", "total_scanned_rows": 1},
    )

    tx_feb = Transaction(
        transaction_date=date(2026, 2, 10),
        description="Feb Txn",
        bank_name="State Bank of India",
        account_identity=ident,
        credit=Decimal("2000.00"),
        source_input_id="inp_feb",
        input_location="Feb_S1_R02",
        statement_id="stmt_feb",
    )
    stmt_feb = BankStatement(
        statement_id="stmt_feb",
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident,
        account_holder="Mr. Rahul Sharma",
        transactions=(tx_feb,),
        metadata={"source_file_stem": "feb_statement", "total_scanned_rows": 1},
    )

    tx_mar = Transaction(
        transaction_date=date(2026, 3, 10),
        description="Mar Txn",
        bank_name="State Bank of India",
        account_identity=ident,
        credit=Decimal("3000.00"),
        source_input_id="inp_mar",
        input_location="Mar_S1_R02",
        statement_id="stmt_mar",
    )
    stmt_mar = BankStatement(
        statement_id="stmt_mar",
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident,
        account_holder="Mr. Rahul Sharma",
        transactions=(tx_mar,),
        metadata={"source_file_stem": "mar_statement", "total_scanned_rows": 1},
    )

    res = consolidate_statements([stmt_jan, stmt_feb, stmt_mar])
    assert len(res.transactions) == 3

    artifact = build_accounts_xlsx_artifact(res)
    wb = openpyxl.load_workbook(io.BytesIO(artifact.content))

    ws_acc = wb["Accounts"]
    acc_rows = list(ws_acc.iter_rows(values_only=True))
    # Must be Header + exactly 1 row for the unique bank account (NOT 3 duplicate account rows!)
    assert len(acc_rows) == 2, f"Expected 1 account row, got {len(acc_rows) - 1} rows in Accounts directory"
    assert acc_rows[1][1] == "Mr. Rahul Sharma"
    assert acc_rows[1][2] == "XXXXXXX6789"
    assert acc_rows[1][3] == "State Bank of India"
    assert acc_rows[1][5] == "TXN-0001 to TXN-0003"

    ws_sum = wb["Processing_Summary"]
    sum_rows = list(ws_sum.iter_rows(values_only=True))
    # Processing Summary must have 3 statement rows + Header + Total row = 5 rows
    assert len(sum_rows) == 5
    # Crucial scope check: Each row must report ONLY its own statement's transaction count (1), NOT account-wide 3!
    assert sum_rows[1][7] == 1, f"Expected 1 transaction in stmt_jan row, got {sum_rows[1][7]}"
    assert sum_rows[2][7] == 1, f"Expected 1 transaction in stmt_feb row, got {sum_rows[2][7]}"
    assert sum_rows[3][7] == 1, f"Expected 1 transaction in stmt_mar row, got {sum_rows[3][7]}"


def test_multi_currency_transactions_xlsx_omits_combined_subtotals() -> None:
    """Multi-currency statements must not combine mixed currencies into a single SUBTOTAL row."""
    import openpyxl

    from sarathi.shakti.bank_statements.consolidator import (
        build_transactions_xlsx_artifact,
        consolidate_statements,
    )

    ident_inr = create_account_identity("State Bank of India", "30123456789")
    tx_inr = Transaction(
        transaction_date=date(2026, 1, 10),
        description="INR Txn",
        bank_name="State Bank of India",
        account_identity=ident_inr,
        credit=Decimal("5000.00"),
        currency="INR",
    )
    stmt_inr = BankStatement(
        statement_id="stmt_inr",
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=ident_inr,
        currency="INR",
        transactions=(tx_inr,),
    )

    ident_usd = create_account_identity("Citibank", "98765432100")
    tx_usd = Transaction(
        transaction_date=date(2026, 1, 15),
        description="USD Txn",
        bank_name="Citibank",
        account_identity=ident_usd,
        credit=Decimal("100.00"),
        currency="USD",
    )
    stmt_usd = BankStatement(
        statement_id="stmt_usd",
        bank_name="Citibank",
        bank_profile="citi",
        account_identity=ident_usd,
        currency="USD",
        transactions=(tx_usd,),
    )

    res = consolidate_statements([stmt_inr, stmt_usd])
    assert len(res.totals_by_currency) == 2
    assert res.total_credit is None  # Scalar total omitted for mixed currencies

    artifact = build_transactions_xlsx_artifact(res)
    wb = openpyxl.load_workbook(io.BytesIO(artifact.content))
    ws = wb["Transactions"]
    rows = list(ws.iter_rows(values_only=True))

    # Header (1), tx_inr (2), tx_usd (3), Totals row (4)
    assert len(rows) == 4
    tot_row = rows[3]
    assert "Mixed Currencies" in str(tot_row[3])
    # Columns G (Debit) and H (Credit) must be None (no single invalid subtotal formula)
    assert tot_row[6] is None
    assert tot_row[7] is None


def test_probable_duplicate_escalates_status_and_consolidate_deduplicates_issues() -> None:
    """Probable duplicates must escalate transaction status to WARNING, and issues must deduplicate across scopes."""
    ident = create_account_identity("SBI", "9988776655")

    tx1 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="UPI Payment To Vendor",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("500.00"),
        sequence_id=1,
    )
    tx2 = Transaction(
        transaction_date=date(2026, 3, 1),
        description="UPI Payment To Vendor",
        bank_name="SBI",
        account_identity=ident,
        debit=Decimal("500.00"),
        sequence_id=2,
    )

    # 1. Deduplication level
    res_dedup = deduplicate_transactions([tx1, tx2])
    assert len(res_dedup.unique_transactions) == 2
    assert res_dedup.unique_transactions[0].status == ValidationStatus.VALID
    assert res_dedup.unique_transactions[1].status == ValidationStatus.WARNING
    assert any(iss.code == "PROBABLE_DUPLICATE_TRANSACTION" for iss in res_dedup.unique_transactions[1].issues)

    # 2. Consolidation level: overall_status must be WARNING
    stmt1 = BankStatement(
        statement_id="stmt1",
        bank_name="SBI",
        bank_profile="sbi",
        account_identity=ident,
        transactions=(tx1,),
    )
    stmt2 = BankStatement(
        statement_id="stmt2",
        bank_name="SBI",
        bank_profile="sbi",
        account_identity=ident,
        transactions=(tx2,),
    )
    res_cons = consolidate_statements([stmt1, stmt2])
    assert res_cons.status == ValidationStatus.WARNING

    # 3. Issue deduplication: duplicate issue codes and messages are not repeated
    issue_keys = [(i.code, i.message) for i in res_cons.issues]
    assert len(issue_keys) == len(set(issue_keys))


def test_near_duplicate_narration_emits_warning():
    """Verify near-duplicate descriptions without references trigger PROBABLE_DUPLICATE_TRANSACTION."""
    ident = create_account_identity("Test Bank", "12345678")
    tx1 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="UPI-SWIGGY-REST-1234",
        bank_name="Test Bank",
        account_identity=ident,
        debit=Decimal("350.00"),
    )
    tx2 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="UPI/SWIGGY/REST/1234",
        bank_name="Test Bank",
        account_identity=ident,
        debit=Decimal("350.00"),
    )

    res = deduplicate_transactions((tx1, tx2))
    assert len(res.unique_transactions) == 2
    # Second transaction should have a warning issue for probable duplicate
    assert any(
        iss.code == "PROBABLE_DUPLICATE_TRANSACTION"
        for t in res.unique_transactions
        for iss in t.issues
    )


def test_multiple_cross_statement_duplicates_preserve_individual_audit_records():
    """Verify that distinct cross-statement duplicates preserve their individual audit issues in consolidation."""
    txs_s1 = [
        Transaction(
            transaction_date=date(2025, 1, i),
            description=f"Payment {i}",
            bank_name="SBI",
            debit=Decimal(str(100 * i)),
            reference_number=f"REF_{i}",
            statement_id="s1",
        )
        for i in range(1, 4)
    ]
    txs_s2 = [
        Transaction(
            transaction_date=date(2025, 1, i),
            description=f"Payment {i}",
            bank_name="SBI",
            debit=Decimal(str(100 * i)),
            reference_number=f"REF_{i}",
            statement_id="s2",
        )
        for i in range(1, 4)
    ]

    ident = create_account_identity("SBI", "123456789012")
    s1 = BankStatement(statement_id="s1", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=tuple(txs_s1))
    s2 = BankStatement(statement_id="s2", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=tuple(txs_s2))

    res = consolidate_statements((s1, s2))
    # All 3 duplicates must be recorded in issues, not collapsed to 1
    dup_issues = [iss for iss in res.issues if iss.code == "CROSS_STATEMENT_DUPLICATE"]
    assert len(dup_issues) == 3


def test_complete_deduplication_reports_zero_successes_in_summary():
    """Verify that when all transactions of statement 2 are removed, Processing Summary reports 0 successes."""
    ident = create_account_identity("SBI", "123456789012")
    tx1 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="Duplicate Row",
        bank_name="SBI",
        debit=Decimal("100.00"),
        reference_number="REF999",
        statement_id="s1",
    )
    tx2 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="Duplicate Row",
        bank_name="SBI",
        debit=Decimal("100.00"),
        reference_number="REF999",
        statement_id="s2",
    )

    s1 = BankStatement(statement_id="s1", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx1,))
    s2 = BankStatement(statement_id="s2", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx2,))

    consolidation = consolidate_statements((s1, s2))
    assert len(consolidation.transactions) == 1

    payload = build_accounts_xlsx_artifact(consolidation)
    wb = openpyxl.load_workbook(io.BytesIO(payload.content))
    ws_sum = wb["Processing_Summary"]

    # Row 2 is Statement 1 (1 success, 0 dupes)
    # Row 3 is Statement 2 (0 successes, 1 dupe)
    s2_succ = ws_sum.cell(row=3, column=8).value
    s2_dups = ws_sum.cell(row=3, column=9).value
    assert s2_succ == 0
    assert s2_dups == 1
    wb.close()


def test_fuzzy_narration_does_not_authorize_proven_duplicate_removal() -> None:
    """Fuzzy narration similarity alone must produce a warning and NEVER remove a transaction as PROVEN_DUPLICATE."""
    ident = create_account_identity("SBI", "123456789012")
    tx1 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="PAYMENT TO ALPHA STORE",
        bank_name="SBI",
        statement_id="s1",
        debit=Decimal("500.00"),
        running_balance=Decimal("4500.00"),
        account_identity=ident,
    )
    tx2 = Transaction(
        transaction_date=date(2025, 1, 10),
        description="PAYMENT TO BETA STORE",
        bank_name="SBI",
        statement_id="s2",
        debit=Decimal("500.00"),
        running_balance=Decimal("4500.00"),
        account_identity=ident,
    )

    s1 = BankStatement(statement_id="s1", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx1,))
    s2 = BankStatement(statement_id="s2", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx2,))

    res = consolidate_statements((s1, s2))
    # Both transactions must be retained; fuzzy similarity cannot delete a transaction
    assert len(res.transactions) == 2
    # Second transaction receives PROBABLE_DUPLICATE warning
    assert any(
        iss.code == "PROBABLE_DUPLICATE_TRANSACTION"
        for t in res.transactions
        for iss in t.issues
    )


def test_deduplication_indices_synchronized_on_enrichment() -> None:
    """Verify that secondary candidate indices are updated when a surviving transaction is enriched with a reference."""
    ident = create_account_identity("SBI", "123456789012")
    # Populate 22 distinct dummy transactions to activate candidate pruning (> 20 candidates)
    dummy_txs = [
        Transaction(
            transaction_date=date(2025, 1, 10),
            description=f"Payment {i}",
            bank_name="SBI",
            statement_id="s1",
            debit=Decimal("10.00"),
            reference_number=f"REF{i:04d}",
            account_identity=ident,
        )
        for i in range(22)
    ]
    # Target transaction without reference
    tx_target = Transaction(
        transaction_date=date(2025, 1, 10),
        description="Enrichment Test",
        bank_name="SBI",
        statement_id="s1",
        debit=Decimal("100.00"),
        running_balance=Decimal("5000.00"),
        account_identity=ident,
    )
    # Matching transaction WITH reference number
    tx_enricher = Transaction(
        transaction_date=date(2025, 1, 10),
        description="Enrichment Test",
        bank_name="SBI",
        statement_id="s2",
        debit=Decimal("100.00"),
        reference_number="REF_NEW_123",
        running_balance=Decimal("5000.00"),
        account_identity=ident,
    )
    # 3rd transaction sharing the enriched reference number
    tx_follower = Transaction(
        transaction_date=date(2025, 1, 10),
        description="Enrichment Test",
        bank_name="SBI",
        statement_id="s3",
        debit=Decimal("100.00"),
        reference_number="REF_NEW_123",
        running_balance=Decimal("5000.00"),
        account_identity=ident,
    )

    all_txs = dummy_txs + [tx_target, tx_enricher, tx_follower]
    res = deduplicate_transactions(all_txs)
    # The 3 matching transactions should all merge into 1 surviving transaction
    # Total unique: 22 dummy + 1 surviving = 23 unique transactions
    assert len(res.unique_transactions) == 23


@pytest.mark.performance
def test_deduplication_candidate_indexing_performance_benchmark():
    """Benchmark: 2,000 candidate transactions sharing account/date/amount with distinct references run in < 0.5s."""
    ident = create_account_identity("SBI", "123456789012")
    n_rows = 2000
    txns = [
        Transaction(
            transaction_date=date(2025, 1, 15),
            description=f"UPI Payment {i}",
            bank_name="SBI",
            debit=Decimal("100.00"),
            reference_number=f"REF{i:06d}",
            account_identity=ident,
        )
        for i in range(n_rows)
    ]

    t_start = time_mod.perf_counter()
    res = deduplicate_transactions(txns)
    elapsed = time_mod.perf_counter() - t_start

    assert len(res.unique_transactions) == n_rows
    # Must complete in well under 0.5 seconds on reference hardware
    assert elapsed < 0.5, f"Deduplication took {elapsed:.3f}s, expected < 0.5s"


@pytest.mark.performance
def test_deduplication_identical_descriptions_distinct_references_benchmark() -> None:
    """Benchmark: 2,000 transactions with identical descriptions and distinct references run in < 0.5s."""
    ident = create_account_identity("SBI", "123456789012")
    n_rows = 2000
    txns = [
        Transaction(
            transaction_date=date(2025, 1, 15),
            description="UPI Payment",  # All identical descriptions!
            bank_name="SBI",
            debit=Decimal("100.00"),
            reference_number=f"REF{i:06d}",
            account_identity=ident,
        )
        for i in range(n_rows)
    ]

    t_start = time_mod.perf_counter()
    res = deduplicate_transactions(txns)
    elapsed = time_mod.perf_counter() - t_start

    assert len(res.unique_transactions) == n_rows
    assert elapsed < 0.5, f"Deduplication took {elapsed:.3f}s, expected < 0.5s"
