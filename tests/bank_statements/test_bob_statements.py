"""Comprehensive Unit & Integration Tests for Bank of Baroda Statements.

Validates:
1. BOB profile identification and in-table identity extraction (Account No, Account Name).
2. 45 transactions extraction with zero-normalization (0.0 -> None in inactive debit/credit).
3. Financial balance reconciliation: Opening Balance (50,829) + Credits (468,711) - Debits (519,540) == Final Balance (0).
4. Option B canonical deliverables and multi-sheet List_of_Accounts.xlsx directory.
"""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from sarathi.sankalpa import ExecutionContext, InputRef, Request, Result
from sarathi.shakti.bank_statements.capability import BankStatementCapability
from sarathi.shakti.bank_statements.detector import detect_bank_statement
from sarathi.shakti.bank_statements.mapper import load_bank_profile_yaml
from sarathi.shakti.native_extraction.readers.spreadsheet import read_xlsx

_BOB_FILE = Path("Input/BOB/10260100019671.xlsx")


@pytest.mark.skipif(not _BOB_FILE.exists(), reason="Input/BOB statement not present")
def test_bob_statement_detection_and_extraction() -> None:
    data = _BOB_FILE.read_bytes()
    doc, _, _ = read_xlsx(data, _BOB_FILE.name)

    det = detect_bank_statement(doc)
    assert det.is_bank_statement is True
    assert det.bank_name == "Bank of Baroda"
    assert det.matched_profile == "bob_excel_fmt1"

    cap = BankStatementCapability()
    in_ref = InputRef(
        input_id="in-bob-01",
        source_path=_BOB_FILE,
        display_name=_BOB_FILE.name,
        size_bytes=len(data),
    )
    req = Request(request_id="req-bob-01", requirement="bank_statement", inputs=(in_ref,))
    ctx = ExecutionContext(run_id="run-bob-01", request_id="req-bob-01", trace_id="tr-1", span_id="sp-1")
    res = cap.execute(req, ctx, Result(data=doc))

    consolidation = res.data
    assert consolidation.total_transactions == 45
    assert len(consolidation.statements) == 1

    stmt = consolidation.statements[0]
    assert stmt.bank_name == "Bank of Baroda"
    assert stmt.account_identity is not None
    assert stmt.account_identity.account_holder == "PRAKASH CHAND JAIN"
    assert stmt.account_identity.masked_account_number.endswith("9671")
    assert len(stmt.transactions) == 45
    assert len(stmt.issues) == 0

    # Financial balance check: 468,711 (credit) and 519,540 (debit)
    assert consolidation.total_credit == Decimal("468711")
    assert consolidation.total_debit == Decimal("519540")
    assert stmt.transactions[-1].running_balance == Decimal("0")

    # Verify clean English narration preservation without legacy font conversion corruption
    assert stmt.transactions[28].description == "Ac xfr from Sol 1026 to 2904"
    assert stmt.transactions[-1].description == "Amt Trfd to RBI-Account DEAF"


@pytest.mark.skipif(not _BOB_FILE.exists(), reason="Input/BOB statement not present")
def test_bob_deliverables_and_formatting() -> None:
    data = _BOB_FILE.read_bytes()
    doc, _, _ = read_xlsx(data, _BOB_FILE.name)

    cap = BankStatementCapability()
    in_ref = InputRef(
        input_id="in-bob-02",
        source_path=_BOB_FILE,
        display_name=_BOB_FILE.name,
        size_bytes=len(data),
    )
    req = Request(request_id="req-bob-02", requirement="bank_statement", inputs=(in_ref,))
    ctx = ExecutionContext(run_id="run-bob-02", request_id="req-bob-02", trace_id="tr-2", span_id="sp-2")
    res = cap.execute(req, ctx, Result(data=doc))
    assert len(res.artifact_payloads) == 3

    # 1. List_of_Accounts.xlsx
    acc_payload = next(a for a in res.artifact_payloads if a.intent.name == "List_of_Accounts.xlsx")
    wb_acc = openpyxl.load_workbook(io.BytesIO(acc_payload.content))
    assert wb_acc.sheetnames == ["Accounts", "Processing_Summary"]

    ws_acc = wb_acc["Accounts"]
    acc_rows = list(ws_acc.iter_rows(values_only=True))
    assert acc_rows[1][1] == "PRAKASH CHAND JAIN"
    assert "9671" in acc_rows[1][2]
    assert acc_rows[1][3] == "Bank of Baroda"

    ws_sum = wb_acc["Processing_Summary"]
    sum_rows = list(ws_sum.iter_rows(values_only=True))
    assert sum_rows[1][1] == _BOB_FILE.stem
    assert sum_rows[1][6] == 45  # Total rows scanned
    assert sum_rows[1][7] == 45  # Successful transactions
    assert sum_rows[1][8] == 0   # Duplicates
    assert sum_rows[1][9] == 0   # Failed/skipped rows
    assert sum_rows[1][10] == "VALID"
    assert sum_rows[1][12] == "ALL_CLEAR"

    # 2. Consolidated_Transactions.xlsx
    txn_payload = next(a for a in res.artifact_payloads if a.intent.name == "Consolidated_Transactions.xlsx")
    wb_txn = openpyxl.load_workbook(io.BytesIO(txn_payload.content))
    ws_txn = wb_txn.active
    # First row is credit 882, debit must be None (zero normalized to blank)
    assert ws_txn.cell(row=2, column=7).value is None
    assert ws_txn.cell(row=2, column=8).value == 882.0

    # 3. Consolidated_Bank_Statement.parquet
    parquet_payload = next(a for a in res.artifact_payloads if a.intent.name == "Consolidated_Bank_Statement.parquet")
    df = pl.read_parquet(io.BytesIO(parquet_payload.content))
    assert len(df) == 45
    assert "transaction_hash" in df.columns
    assert "account_fingerprint" in df.columns


def test_bob_profile_file_format_metadata() -> None:
    profile_path = Path("src/sarathi/data/banks/bob_excel_fmt1.yaml")
    assert profile_path.exists()
    cfg = load_bank_profile_yaml(profile_path)
    assert cfg["profile_id"] == "bob_excel_fmt1"
    assert cfg["bank_name"] == "Bank of Baroda"
    assert cfg["file_format"]["declared_extension"] == ".xlsx"
