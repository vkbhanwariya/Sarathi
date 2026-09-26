"""Edge Case Audit Regression & Performance Suite for Bank Statements in Sarathi.

Covers:
1. Worksheet names (e.g. 'Transactions', 'Statement') not becoming fake account numbers.
2. Distinct banks with identical masked tails (XXXX1234) never merged.
3. Same account with and without IFSC properly grouped and deduplicated.
4. Multi-page continuation table without repeated headers extracts all pages.
5. Narration with hyphens ('UPI---SHOP') not classified as noise.
6. Near-duplicate descriptions (OCR/spacing noise) emitting PROBABLE_DUPLICATE_TRANSACTION.
7. Multiple cross-statement duplicates preserving distinct audit records in seen_issues.
8. Cross-statement balance discontinuity and missing statement period detection.
9. Outcome aggregation by source_input_id with invalid date error propagation.
10. Complete deduplication reporting 0 successful transactions in Processing Summary.
11. Reverse-order same-day transactions avoiding false balance continuity warnings.
12. Parquet eod_balance exporting true end-of-day balance (not intraday).
13. High-cardinality candidate deduplication performance benchmark (< 0.5s for 2,000 rows).
"""

from __future__ import annotations

import datetime
import io
import time
from decimal import Decimal
from pathlib import Path

import openpyxl
import polars as pl

from sarathi.sankalpa import (
    CanonicalDocument,
    ExecutionContext,
    InputRef,
    PageData,
    Request,
    Result,
    TableData,
)
from sarathi.shakti.bank_statements.capability import BankStatementCapability
from sarathi.shakti.bank_statements.consolidator import (
    _account_group_key,
    build_accounts_xlsx_artifact,
    build_parquet_artifact,
    consolidate_statements,
)
from sarathi.shakti.bank_statements.deduplicator import deduplicate_transactions
from sarathi.shakti.bank_statements.models import (
    AccountIdentity,
    BankStatement,
    Transaction,
    create_account_identity,
)
from sarathi.shakti.bank_statements.row_classifier import RowType, classify_row
from sarathi.shakti.bank_statements.validator import validate_statement_balances


def test_worksheet_name_does_not_become_fake_account():
    """Verify that sheet names like 'Transactions' or 'Statement' are rejected as account numbers."""
    cap = BankStatementCapability()
    req = Request(
        request_id="req-1",
        requirement="bank_statements",
        inputs=(InputRef("inp1", Path("test.xlsx"), "test.xlsx", 100),),
    )
    ctx = ExecutionContext("run-1", "req-1", "t1", "s1")

    # Document with a table named "Transactions"
    doc = CanonicalDocument(
        document_id="doc1",
        source_input_id="inp1",
        text="",
        tables=(
            TableData(
                headers=("Date", "Description", "Debit", "Credit", "Balance"),
                rows=(
                    ("01/01/2025", "Salary", "", "50000", "50000"),
                ),
                name="Transactions",
            ),
        ),
    )

    res = cap.execute(req, ctx, prior_result=Result(data=doc))
    consolidation = res.data
    assert len(consolidation.transactions) == 1
    stmt = consolidation.statements[0]
    ident = stmt.account_identity
    # Account identity must NOT have account_number "Transactions"
    if ident and ident.masked_account_number:
        assert "Transactions" not in ident.masked_account_number


def test_distinct_banks_with_identical_masked_tails_not_merged():
    """Verify accounts from different banks with identical masked tails are not merged."""
    stmt1 = BankStatement(
        statement_id="stmt_sbi",
        bank_name="State Bank of India",
        bank_profile="sbi",
        account_identity=AccountIdentity(
            bank_name="State Bank of India",
            masked_account_number="XXXX1234",
            identity_strength="WEAK",
        ),
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 1, 1),
                description="Tx 1",
                bank_name="State Bank of India",
                debit=Decimal("100"),
                statement_id="stmt_sbi",
            ),
        ),
    )

    stmt2 = BankStatement(
        statement_id="stmt_hdfc",
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=AccountIdentity(
            bank_name="HDFC Bank",
            masked_account_number="XXXX1234",
            identity_strength="WEAK",
        ),
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 1, 2),
                description="Tx 2",
                bank_name="HDFC Bank",
                debit=Decimal("200"),
                statement_id="stmt_hdfc",
            ),
        ),
    )

    # Grouping keys must differ
    k1 = _account_group_key(stmt1)
    k2 = _account_group_key(stmt2)
    assert k1 != k2

    consolidation = consolidate_statements((stmt1, stmt2))
    xlsx_payload = build_accounts_xlsx_artifact(consolidation)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_payload.content))
    ws_acc = wb["Accounts"]
    # Master Account Directory must have 2 distinct rows (row 2 and row 3)
    assert ws_acc.max_row == 3
    wb.close()


def test_same_account_with_and_without_ifsc_groups_and_deduplicates():
    """Verify that statements for the same verified account group together even if IFSC is missing in one."""
    ident_with_ifsc = create_account_identity(
        bank_name="HDFC Bank",
        raw_account_number="50100123456789",
        account_holder="Alice",
        ifsc="HDFC0001234",
    )
    ident_without_ifsc = create_account_identity(
        bank_name="HDFC Bank",
        raw_account_number="50100123456789",
        account_holder="Alice",
        ifsc=None,
    )

    stmt1 = BankStatement(
        statement_id="stmt_month1",
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident_with_ifsc,
        ifsc="HDFC0001234",
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 1, 15),
                description="Shared Tx",
                bank_name="HDFC Bank",
                debit=Decimal("500"),
                reference_number="REF12345",
                statement_id="stmt_month1",
            ),
        ),
    )
    stmt2 = BankStatement(
        statement_id="stmt_month2",
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident_without_ifsc,
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 1, 15),
                description="Shared Tx",
                bank_name="HDFC Bank",
                debit=Decimal("500"),
                reference_number="REF12345",
                statement_id="stmt_month2",
            ),
        ),
    )

    k1 = _account_group_key(stmt1)
    k2 = _account_group_key(stmt2)
    assert k1 == k2

    consolidation = consolidate_statements((stmt1, stmt2))
    # Duplicate transaction must be eliminated across statements
    assert len(consolidation.transactions) == 1


def test_multipage_continuation_table_without_repeated_headers():
    """Verify that page 2 continuation tables without repeated headers are successfully extracted."""
    cap = BankStatementCapability()
    req = Request(
        request_id="req-multi",
        requirement="bank_statements",
        inputs=(InputRef("inp_multi", Path("stmt.pdf"), "stmt.pdf", 100),),
    )
    ctx = ExecutionContext("run-multi", "req-multi", "t1", "s1")

    page1_table = TableData(
        headers=("Date", "Description", "Debit", "Credit", "Balance"),
        rows=(
            ("01/01/2025", "Tx Page 1", "100", "", "900"),
        ),
    )
    page2_table = TableData(
        headers=(),
        rows=(
            ("02/01/2025", "Tx Page 2", "", "200", "1100"),
        ),
    )

    doc = CanonicalDocument(
        document_id="doc_multi",
        source_input_id="inp_multi",
        text="",
        pages=(
            PageData(page_number=1, tables=(page1_table,)),
            PageData(page_number=2, tables=(page2_table,)),
        ),
    )

    res = cap.execute(req, ctx, prior_result=Result(data=doc))
    consolidation = res.data
    # Both page 1 and page 2 transactions must be extracted
    assert len(consolidation.transactions) == 2
    assert consolidation.transactions[0].description == "Tx Page 1"
    assert consolidation.transactions[1].description == "Tx Page 2"


def test_narration_with_hyphens_not_classified_as_noise():
    """Verify that a transaction narration like 'UPI---SHOP' is classified as TRANSACTION, not NOISE."""
    row = ["01/01/2025", "UPI---SHOP", "500.00", "", "1000.00"]
    r_type = classify_row(row, date_col_idx=0, amount_col_indices=[2, 3, 4])
    assert r_type == RowType.TRANSACTION


def test_near_duplicate_narration_emits_warning():
    """Verify near-duplicate descriptions without references trigger PROBABLE_DUPLICATE_TRANSACTION."""
    ident = create_account_identity("Test Bank", "12345678")
    tx1 = Transaction(
        transaction_date=datetime.date(2025, 1, 10),
        description="UPI-SWIGGY-REST-1234",
        bank_name="Test Bank",
        account_identity=ident,
        debit=Decimal("350.00"),
    )
    tx2 = Transaction(
        transaction_date=datetime.date(2025, 1, 10),
        description="UPI/SWIGGY/REST/1234",
        bank_name="Test Bank",
        account_identity=ident,
        debit=Decimal("350.00"),
    )

    res = deduplicate_transactions((tx1, tx2))
    assert len(res.unique_transactions) == 2
    # Second transaction should have a warning issue for probable duplicate
    assert any(
        iss.code == "PROBABLE_DUPLICATE_TRANSACTION"
        for t in res.unique_transactions
        for iss in t.issues
    )


def test_multiple_cross_statement_duplicates_preserve_individual_audit_records():
    """Verify that distinct cross-statement duplicates preserve their individual audit issues in consolidation."""
    txs_s1 = [
        Transaction(
            transaction_date=datetime.date(2025, 1, i),
            description=f"Payment {i}",
            bank_name="SBI",
            debit=Decimal(str(100 * i)),
            reference_number=f"REF_{i}",
            statement_id="s1",
        )
        for i in range(1, 4)
    ]
    txs_s2 = [
        Transaction(
            transaction_date=datetime.date(2025, 1, i),
            description=f"Payment {i}",
            bank_name="SBI",
            debit=Decimal(str(100 * i)),
            reference_number=f"REF_{i}",
            statement_id="s2",
        )
        for i in range(1, 4)
    ]

    ident = create_account_identity("SBI", "123456789012")
    s1 = BankStatement(statement_id="s1", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=tuple(txs_s1))
    s2 = BankStatement(statement_id="s2", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=tuple(txs_s2))

    res = consolidate_statements((s1, s2))
    # All 3 transactions were duplicated across statements
    dup_issues = [iss for iss in res.issues if iss.code == "CROSS_STATEMENT_DUPLICATE"]
    assert len(dup_issues) == 3


def test_cross_statement_balance_discontinuity_and_period_gap_detection():
    """Verify that cross-statement balance gap and missing month gaps are detected."""
    ident = create_account_identity("HDFC Bank", "50100987654321")

    # Statement 1 ends on Jan 31 with closing balance 1000
    s1 = BankStatement(
        statement_id="s1",
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident,
        statement_period_start=datetime.date(2025, 1, 1),
        statement_period_end=datetime.date(2025, 1, 31),
        opening_balance=Decimal("500"),
        closing_balance=Decimal("1000"),
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 1, 15),
                description="Jan Tx",
                bank_name="HDFC Bank",
                credit=Decimal("500"),
                running_balance=Decimal("1000"),
            ),
        ),
    )

    # Statement 2 starts on March 15 (gap > 35 days) with opening balance 1200 (discontinuity != 1000)
    s2 = BankStatement(
        statement_id="s2",
        bank_name="HDFC Bank",
        bank_profile="hdfc",
        account_identity=ident,
        statement_period_start=datetime.date(2025, 3, 15),
        statement_period_end=datetime.date(2025, 3, 31),
        opening_balance=Decimal("1200"),
        closing_balance=Decimal("1500"),
        transactions=(
            Transaction(
                transaction_date=datetime.date(2025, 3, 20),
                description="Mar Tx",
                bank_name="HDFC Bank",
                credit=Decimal("300"),
                running_balance=Decimal("1500"),
            ),
        ),
    )

    res = consolidate_statements((s1, s2))
    assert any(iss.code == "CROSS_STATEMENT_BALANCE_DISCONTINUITY" for iss in res.issues)
    assert any(iss.code == "MISSING_STATEMENT_PERIOD" for iss in res.issues)


def test_complete_deduplication_reports_zero_successes_in_summary():
    """Verify that when all transactions of statement 2 are removed, Processing Summary reports 0 successes."""
    ident = create_account_identity("SBI", "123456789012")
    tx1 = Transaction(
        transaction_date=datetime.date(2025, 1, 10),
        description="Duplicate Row",
        bank_name="SBI",
        debit=Decimal("100.00"),
        reference_number="REF999",
        statement_id="s1",
    )
    tx2 = Transaction(
        transaction_date=datetime.date(2025, 1, 10),
        description="Duplicate Row",
        bank_name="SBI",
        debit=Decimal("100.00"),
        reference_number="REF999",
        statement_id="s2",
    )

    s1 = BankStatement(statement_id="s1", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx1,))
    s2 = BankStatement(statement_id="s2", bank_name="SBI", bank_profile="sbi", account_identity=ident, transactions=(tx2,))

    consolidation = consolidate_statements((s1, s2))
    assert len(consolidation.transactions) == 1

    payload = build_accounts_xlsx_artifact(consolidation)
    wb = openpyxl.load_workbook(io.BytesIO(payload.content))
    ws_sum = wb["Processing_Summary"]

    # Row 2 is Statement 1 (1 success, 0 dupes)
    # Row 3 is Statement 2 (0 successes, 1 dupe)
    s2_succ = ws_sum.cell(row=3, column=8).value
    s2_dups = ws_sum.cell(row=3, column=9).value
    assert s2_succ == 0
    assert s2_dups == 1
    wb.close()


def test_reverse_order_same_day_transactions_no_false_discontinuity():
    """Verify two newest-first transactions on the same date with correct balances do not emit false balance errors."""
    # Newest-first:
    # Row 1 (later in day): debit 5, balance 85
    # Row 2 (earlier in day): initial balance was 90
    t1 = Transaction(
        transaction_date=datetime.date(2025, 1, 15),
        description="Later Tx",
        bank_name="Test Bank",
        debit=Decimal("5.00"),
        running_balance=Decimal("85.00"),
    )
    t2 = Transaction(
        transaction_date=datetime.date(2025, 1, 15),
        description="Earlier Tx",
        bank_name="Test Bank",
        credit=Decimal("10.00"),
        running_balance=Decimal("90.00"),
    )

    stmt = BankStatement(
        statement_id="stmt_rev",
        bank_name="Test Bank",
        bank_profile="common",
        opening_balance=Decimal("80.00"),
        closing_balance=Decimal("85.00"),
        transactions=(t1, t2),
    )

    val_stmt = validate_statement_balances(stmt)
    assert not any(iss.code == "RUNNING_BALANCE_DISCONTINUITY" for iss in val_stmt.issues)


def test_parquet_eod_balance_is_end_of_day_not_intraday():
    """Verify Parquet exports true EOD balance for same-day transactions."""
    t1 = Transaction(
        transaction_date=datetime.date(2025, 1, 15),
        transaction_time=datetime.time(10, 0),
        description="Morning Tx",
        bank_name="Test Bank",
        credit=Decimal("10.00"),
        running_balance=Decimal("90.00"),
        statement_id="s1",
    )
    t2 = Transaction(
        transaction_date=datetime.date(2025, 1, 15),
        transaction_time=datetime.time(14, 0),
        description="Afternoon Tx",
        bank_name="Test Bank",
        debit=Decimal("5.00"),
        running_balance=Decimal("85.00"),
        statement_id="s1",
    )

    s1 = BankStatement(
        statement_id="s1",
        bank_name="Test Bank",
        bank_profile="common",
        statement_period_start=datetime.date(2025, 1, 1),
        statement_period_end=datetime.date(2025, 1, 31),
        opening_balance=Decimal("80.00"),
        closing_balance=Decimal("85.00"),
        transactions=(t1, t2),
    )

    consolidation = consolidate_statements((s1,))
    payload = build_parquet_artifact(consolidation)
    df = pl.read_parquet(io.BytesIO(payload.content))

    eod_values = df["eod_balance"].to_list()
    # Both transactions on Jan 15 must have EOD balance 85.00
    assert eod_values[0] == Decimal("85.00")
    assert eod_values[1] == Decimal("85.00")


def test_outcome_aggregation_by_source_input_id():
    """Verify that multiple documents with distinct outcomes report correct SUCCESS/FAILED by input ID."""
    cap = BankStatementCapability()
    req = Request(
        request_id="req-outcomes",
        requirement="bank_statements",
        inputs=(
            InputRef("inp_valid", Path("valid.xlsx"), "valid.xlsx", 100),
            InputRef("inp_invalid", Path("invalid.xlsx"), "invalid.xlsx", 100),
        ),
    )
    ctx = ExecutionContext("run-outcomes", "req-outcomes", "t1", "s1")

    doc_valid = CanonicalDocument(
        document_id="doc_valid",
        source_input_id="inp_valid",
        text="",
        tables=(
            TableData(
                headers=("Date", "Description", "Debit", "Credit", "Balance"),
                rows=(
                    ("01/01/2025", "Valid Tx", "100", "", "900"),
                ),
            ),
        ),
    )

    doc_invalid = CanonicalDocument(
        document_id="doc_invalid",
        source_input_id="inp_invalid",
        text="",
        tables=(
            TableData(
                headers=("Date", "Description", "Debit", "Credit", "Balance"),
                rows=(
                    ("INVALID_DATE_TEXT", "Bad Tx", "100", "", "900"),
                ),
            ),
        ),
    )

    res = cap.execute(req, ctx, prior_result=Result(data=(doc_valid, doc_invalid)))
    input_outcomes = res.metadata.get("input_outcomes", {})
    assert input_outcomes.get("inp_valid") == "SUCCESS"
    assert input_outcomes.get("inp_invalid") == "FAILED"


def test_deduplication_candidate_indexing_performance_benchmark():
    """Benchmark: 2,000 candidate transactions sharing account/date/amount with distinct references run in < 0.5s."""
    ident = create_account_identity("SBI", "123456789012")
    n_rows = 2000
    txns = [
        Transaction(
            transaction_date=datetime.date(2025, 1, 15),
            description=f"UPI Payment {i}",
            bank_name="SBI",
            debit=Decimal("100.00"),
            reference_number=f"REF{i:06d}",
            account_identity=ident,
        )
        for i in range(n_rows)
    ]

    t_start = time.perf_counter()
    res = deduplicate_transactions(txns)
    elapsed = time.perf_counter() - t_start

    assert len(res.unique_transactions) == n_rows
    # Must complete in well under 0.5 seconds on reference hardware
    assert elapsed < 0.5, f"Deduplication took {elapsed:.3f}s, expected < 0.5s"
