"""Tests for In-Place Excel (.xlsx) Translation Transcoder and Table Exporter."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from sarathi.sankalpa import (
    CanonicalDocument,
    ExecutionContext,
    InputRef,
    Request,
    Result,
    TableData,
    WarningRecord,
)
from sarathi.shakti.translation.capability import TranslationCapability
from sarathi.shakti.translation.engine import TranslatorBackend
from sarathi.shakti.translation.xlsx_transformer import (
    build_xlsx_from_tables,
    transform_xlsx_translation_artifact,
)


def test_transform_xlsx_preserves_formulas_numbers_and_styles(tmp_path: Path) -> None:
    """Verify in-place XLSX translation translates text while strictly preserving formulas, numbers, and styling."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"

    header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    ws["A1"] = "Revenue Statement"
    ws["A1"].font = header_font
    ws["A1"].fill = header_fill

    ws["A2"] = "Total Income"
    ws["B2"] = 50000
    ws["B2"].number_format = "₹#,##0"

    ws["A3"] = "Tax Deducted"
    ws["B3"] = 9000

    ws["A4"] = "Net Profit"
    ws["B4"] = "=B2-B3"  # Formula that must NEVER be translated or corrupted

    buf = io.BytesIO()
    wb.save(buf)
    input_bytes = buf.getvalue()

    translation_dict = {
        "Financial Summary": "वित्तीय सारांश",
        "Revenue Statement": "राजस्व विवरण",
        "Total Income": "कुल आय",
        "Tax Deducted": "कटौती किया गया कर",
        "Net Profit": "शुद्ध लाभ",
    }

    def mock_translate_fn(batch: list[str]) -> list[str]:
        return [translation_dict.get(s, f"TR:{s}") for s in batch]

    payload = transform_xlsx_translation_artifact(
        input_bytes=input_bytes,
        translate_fn=mock_translate_fn,
        filename="Translated_Summary.xlsx",
        is_hindi_target=True,
    )

    assert payload.intent.name == "Translated_Summary.xlsx"
    assert (
        payload.intent.media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    # Re-open the generated workbook and inspect cells
    res_wb = openpyxl.load_workbook(io.BytesIO(payload.content), data_only=False)
    res_ws = res_wb.active

    # 1. Sheet title preserved by default to protect cross-sheet formula references
    assert res_ws.title == "Financial Summary"

    # 2. Text translated with font adapted to Nirmala UI while preserving bold and color
    assert res_ws["A1"].value == "राजस्व विवरण"
    assert res_ws["A1"].font.name == "Nirmala UI"
    assert res_ws["A1"].font.bold is True
    assert res_ws["A1"].font.color.rgb == "00FFFFFF" or res_ws["A1"].font.color.value == "FFFFFF"

    assert res_ws["A2"].value == "कुल आय"
    assert res_ws["A2"].font.name == "Nirmala UI"

    # 3. Numeric cells untouched
    assert res_ws["B2"].value == 50000
    assert res_ws["B2"].number_format == "₹#,##0"

    assert res_ws["A3"].value == "कटौती किया गया कर"
    assert res_ws["B3"].value == 9000

    assert res_ws["A4"].value == "शुद्ध लाभ"

    # 4. Critical: formula cell untouched
    assert res_ws["B4"].value == "=B2-B3"


def test_transform_xlsx_preserves_cross_sheet_formulas() -> None:
    """Verify cross-sheet formulas like =Data!A1 remain valid after translation."""
    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Data"
    ws_data["A1"] = 500
    ws_data["A2"] = "Source Data Label"

    ws_summary = wb.create_sheet(title="Summary")
    ws_summary["A1"] = "Total"
    ws_summary["B1"] = "=Data!A1"

    buf = io.BytesIO()
    wb.save(buf)
    input_bytes = buf.getvalue()

    def mock_trans(batch: list[str]) -> list[str]:
        return [f"TR_{s}" for s in batch]

    payload = transform_xlsx_translation_artifact(
        input_bytes=input_bytes,
        translate_fn=mock_trans,
        filename="Translated_CrossSheet.xlsx",
        is_hindi_target=True,
    )

    res_wb = openpyxl.load_workbook(io.BytesIO(payload.content), data_only=False)
    # Sheet names must be preserved so cross-sheet references are not orphaned into #REF!
    assert "Data" in res_wb.sheetnames
    assert "Summary" in res_wb.sheetnames
    assert res_wb["Summary"]["B1"].value == "=Data!A1"
    assert res_wb["Summary"]["A1"].value == "TR_Total"
    assert res_wb["Data"]["A2"].value == "TR_Source Data Label"
    assert res_wb["Data"]["A1"].value == 500


def test_build_xlsx_from_tables() -> None:
    """Verify synthesis of XLSX workbook from TableData instances."""
    t1 = TableData(
        name="Case Details",
        headers=("Case Number", "Court Name", "Status"),
        rows=(
            ("CR-124/2023", "High Court of Delhi", "Pending"),
            ("BA-55/2024", "Sessions Court", "Disposed"),
        ),
    )
    t2 = TableData(
        name="Amounts",
        headers=("Description", "Amount in INR"),
        rows=(
            ("Personal Bond", "50000"),
            ("Surety Amount", "100000.50"),
        ),
    )

    payload = build_xlsx_from_tables([t1, t2], filename="Translated_Tables.xlsx", is_hindi_target=True)
    assert payload.intent.name == "Translated_Tables.xlsx"

    wb = openpyxl.load_workbook(io.BytesIO(payload.content))
    assert len(wb.sheetnames) == 2
    assert wb.sheetnames[0] == "Case Details"
    assert wb.sheetnames[1] == "Amounts"

    ws1 = wb["Case Details"]
    assert ws1.cell(row=1, column=1).value == "Case Number"
    assert ws1.cell(row=2, column=2).value == "High Court of Delhi"

    ws2 = wb["Amounts"]
    assert ws2.cell(row=1, column=2).value == "Amount in INR"
    # Auto-numeric conversion verified
    assert ws2.cell(row=2, column=2).value == 50000
    assert ws2.cell(row=3, column=2).value == 100000.50


def test_translation_capability_xlsx_e2e(tmp_path: Path) -> None:
    """Verify TranslationCapability produces an in-place translated .xlsx alongside DOCX and TXT."""
    # 1. Create source XLSX on disk
    src_xlsx = tmp_path / "orders.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws["A1"] = "Order No"
    ws["B1"] = "Party"
    ws["A2"] = "1"
    ws["B2"] = "State of Rajasthan"
    wb.save(src_xlsx)

    class MockBackend(TranslatorBackend):
        def translate_sentences(
            self, sentences: list[str], direction, execution_binding=None, engine: str = "krutrim", **kwargs
        ) -> tuple[list[str], str]:
            return [f"TR:{s}" for s in sentences], "mock_krutrim"

    cap = TranslationCapability(backend=MockBackend())
    doc = CanonicalDocument(
        document_id="doc-orders",
        source_input_id="inp-orders",
        text="Order No Party",
    )
    req = Request(
        request_id="req-1",
        requirement="translation",
        inputs=(InputRef("inp-orders", src_xlsx, "orders.xlsx", src_xlsx.stat().st_size),),
        custom_options={"direction": "en_hi"},
    )
    ctx = ExecutionContext("req-1", "translation", "trace-1", "span-1")

    res = cap.execute(req, ctx, prior_result=Result(data=doc))
    assert res.data is not None

    # Verify artifacts include .txt, .docx, AND .xlsx
    artifact_names = [p.intent.name for p in res.artifact_payloads]
    assert any(n.endswith(".txt") for n in artifact_names)
    assert any(n.endswith(".docx") for n in artifact_names)
    assert any(n.endswith(".xlsx") for n in artifact_names)

    xlsx_art = next(p for p in res.artifact_payloads if p.intent.name.endswith(".xlsx"))
    res_wb = openpyxl.load_workbook(io.BytesIO(xlsx_art.content))
    res_ws = res_wb.active
    assert "TR:" in str(res_ws["B2"].value)


def test_translation_capability_csv_tabular_e2e(tmp_path: Path) -> None:
    """Verify TranslationCapability produces an XLSX workbook when translating a CSV/tabular document."""
    src_csv = tmp_path / "data.csv"
    src_csv.write_text("Item,Quantity,Rate\nWheat,100,2500\nRice,50,4000\n", encoding="utf-8")

    class MockBackend(TranslatorBackend):
        def translate_sentences(
            self, sentences: list[str], direction, execution_binding=None, engine: str = "krutrim", **kwargs
        ) -> tuple[list[str], str]:
            return [f"TR:{s}" for s in sentences], "mock_krutrim"

    cap = TranslationCapability(backend=MockBackend())
    table = TableData(
        name="Items",
        headers=("Item", "Quantity", "Rate"),
        rows=(("Wheat", "100", "2500"), ("Rice", "50", "4000")),
    )
    doc = CanonicalDocument(
        document_id="doc-csv",
        source_input_id="inp-csv",
        detected_type="csv",
        text="Item Quantity Rate Wheat 100 2500 Rice 50 4000",
        tables=[table],
    )
    req = Request(
        request_id="req-csv",
        requirement="translation",
        inputs=(InputRef("inp-csv", src_csv, "data.csv", src_csv.stat().st_size),),
        custom_options={"direction": "en_hi"},
    )
    ctx = ExecutionContext("req-csv", "translation", "trace-csv", "span-csv")

    res = cap.execute(req, ctx, prior_result=Result(data=doc))
    assert res.data is not None

    artifact_names = [p.intent.name for p in res.artifact_payloads]
    assert any(n.endswith(".xlsx") for n in artifact_names)
    xlsx_art = next(p for p in res.artifact_payloads if p.intent.name.endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_art.content))
    ws = wb.active
    assert "TR:" in str(ws.cell(row=2, column=1).value)


def test_transform_xlsm_preserves_vba_and_media_type() -> None:
    """Verify that macro-enabled .xlsm files preserve VBA macros and export as .xlsm."""
    import zipfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Revenue"
    ws["B1"] = 1000

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    # Inject an inert dummy vbaProject.bin into the package to simulate an .xlsm workbook
    xlsm_buf = io.BytesIO()
    src_io = io.BytesIO(buf.getvalue())
    with zipfile.ZipFile(src_io, "r") as src_zip:
        with zipfile.ZipFile(xlsm_buf, "w") as dst_zip:
            for item in src_zip.infolist():
                data = src_zip.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                    )
                dst_zip.writestr(item, data)
            dst_zip.writestr("xl/vbaProject.bin", b"VBA_DUMMY_CODE_PACKAGE")

    input_xlsm = xlsm_buf.getvalue()
    warns: list[WarningRecord] = []

    payload = transform_xlsx_translation_artifact(
        input_bytes=input_xlsm,
        translate_fn=lambda b: [f"HI_{x}" for x in b],
        filename="Financial_Report.xlsm",
        warnings=warns,
    )

    assert payload.intent.name.endswith(".xlsm")
    assert payload.intent.media_type == "application/vnd.ms-excel.sheet.macroEnabled.12"

    res_io = io.BytesIO(payload.content)
    with zipfile.ZipFile(res_io, "r") as zf:
        assert "xl/vbaProject.bin" in zf.namelist()
        assert zf.read("xl/vbaProject.bin") == b"VBA_DUMMY_CODE_PACKAGE"


def test_transform_xlsx_synchronizes_table_column_names() -> None:
    """Verify openpyxl TableColumn names are updated in sync with translated header cells."""
    from openpyxl.worksheet.table import Table, TableColumn

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EntriesSheet"
    ws.append(["Category", "Amount"])
    ws.append(["Groceries", 500])
    tab = Table(displayName="EntriesTable", ref="A1:B2")
    tab.tableColumns = [TableColumn(id=1, name="Category"), TableColumn(id=2, name="Amount")]
    ws.add_table(tab)

    buf = io.BytesIO()
    wb.save(buf)
    input_bytes = buf.getvalue()

    trans_map = {"Category": "श्रेणी", "Amount": "राशि", "Groceries": "किराना"}
    payload = transform_xlsx_translation_artifact(
        input_bytes=input_bytes,
        translate_fn=lambda b: [trans_map.get(x, x) for x in b],
        filename="Table_Doc.xlsx",
        is_hindi_target=True,
    )

    out_io = io.BytesIO(payload.content)
    res_wb = openpyxl.load_workbook(out_io)
    res_ws = res_wb.active
    res_tab = list(res_ws.tables.values())[0]

    assert [c.name for c in res_tab.tableColumns] == ["श्रेणी", "राशि"]
    assert res_ws["A1"].value == "श्रेणी"
    assert res_ws["B1"].value == "राशि"
    res_wb.close()
